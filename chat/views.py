import ast
import base64
import csv
import difflib
import hashlib
import io
import json
import os
import re
import secrets
import sqlite3
import subprocess
import time
import uuid
import zipfile
import tempfile
from pathlib import PurePosixPath
from urllib.parse import quote, urlencode, urlparse

import requests

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    import fitz
except ImportError:
    fitz = None

try:
    from faster_whisper import WhisperModel
except ImportError:
    WhisperModel = None

_whisper_model = None

from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.conf import settings
from django.db.models import Count, Min, Q
from django.http import HttpResponse, JsonResponse, StreamingHttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_GET, require_http_methods, require_POST
from django.views.decorators.csrf import csrf_exempt

from .models import (
    ChatMessage,
    ChatSession,
    ConversationRevision,
    KnowledgeChunk,
    KnowledgeDocument,
    LocalModelConfig,
    McpConnector,
    McpToolCall,
    AgentTask,
    AgentTeam,
    SandboxPolicy,
    AgentJob,
    AiEvent,
    AuditEvent,
    UserOllamaSettings,
    PrivacyPreference,
    Workspace,
    WorkspaceMembership,
    WorkspacePolicy,
    EnterpriseIdentityConfig,
    DirectoryProvisioningEvent,
    EvaluationTask,
    EvaluationRun,
    EvaluationScore,
    EvaluationRegressionSuite,
    SecretVaultItem,
    ExtensionPackage,
    ExtensionInstall,
)
from .quality import analyze_code_quality
from .sandbox import run_sandboxed_code
from .security import scan_files
from .devops import analyze_logs, generate_artifact
from .documentation import generate_documentation
from .dependencies import analyze_dependencies
from .browser_testing import analyze_browser_report, generate_playwright_test
from .model_router import route_model
from .devcontainer import generate_devcontainer
from .deployment import generate_deployment_kit, validate_deployment_environment
from .incident import analyze_incident
from .architecture import analyze_architecture
from .cross_repository import analyze_cross_repository
from .vault import decrypt_blob, decrypt_secret, encrypt_blob, encrypt_secret, mask_json, redact_sensitive_text
from .api_contract import analyze_api_contract
from .provenance import generate_provenance, verify_provenance
from .review import review_gate


OLLAMA_BASE_URL = "http://127.0.0.1:11434"
OLLAMA_URL = f"{OLLAMA_BASE_URL}/api/generate"
DEFAULT_MODEL = "qwen2.5-coder:1.5b"
VISION_MODEL = os.environ.get("OLLAMA_VISION_MODEL", "llava:latest")
MAX_PROJECT_FILES = 100
MAX_FILE_BYTES = 1_000_000
MAX_IMAGE_BYTES = 5_000_000
CHUNK_SIZE = 2_000
MAX_EXECUTION_CHARS = 20_000
MAX_TEST_CHARS = 20_000
MAX_GIT_OUTPUT_CHARS = 50_000
MAX_BACKUP_BYTES = 25_000_000

EXTENSION_CATALOG = [
    {
        "slug": "mcp-project-search",
        "name": "Project Search Connector",
        "version": "1.1.0",
        "description": "Read-only project search for local MCP workflows.",
        "permissions": ["project.read"],
        "manifest": {"connector_type": "documentation", "tool": "project.search"},
    },
    {
        "slug": "agent-reviewer",
        "name": "Agent Review Assistant",
        "version": "1.0.0",
        "description": "Adds a guarded review checkpoint to agent tasks.",
        "permissions": ["agent.read", "agent.approve"],
        "manifest": {"agent_role": "reviewer", "approval_required": True},
    },
    {
        "slug": "docs-indexer",
        "name": "Documentation Indexer",
        "version": "1.0.0",
        "description": "Indexes local documentation for project-aware answers.",
        "permissions": ["project.read", "index.write"],
        "manifest": {"connector_type": "documentation", "index": "local"},
    },
]


def _agent_plan(goal):
    plan = [
        {
            "title": "Inspect project context",
            "description": "Read indexed files and relevant project history.",
            "kind": "read",
            "requires_approval": False,
            "approved": True,
            "completed": False,
        },
        {
            "title": "Prepare file changes",
            "description": "Propose edits for the files related to this task.",
            "kind": "file_write",
            "requires_approval": True,
            "approved": False,
            "completed": False,
        },
        {
            "title": "Run guarded tests",
            "description": "Execute the project test command inside the guarded workspace.",
            "kind": "execute",
            "requires_approval": True,
            "approved": False,
            "completed": False,
        },
        {
            "title": "Review Git changes",
            "description": "Inspect the resulting diff and changed files.",
            "kind": "git_read",
            "requires_approval": False,
            "approved": True,
            "completed": False,
        },
        {
            "title": "Create a Git commit",
            "description": "Prepare a commit only after explicit user approval.",
            "kind": "git_write",
            "requires_approval": True,
            "approved": False,
            "completed": False,
        },
    ]
    return plan


def _agent_payload(task):
    completed = sum(1 for step in task.plan if step.get("completed"))
    return {
        "id": task.id,
        "title": task.title,
        "goal": task.goal,
        "status": task.status,
        "plan": task.plan,
        "current_step": task.current_step,
        "progress": completed,
        "total_steps": len(task.plan),
        "result": task.result,
        "control_state": task.control_state,
        "logs": task.logs,
        "retry_count": task.retry_count,
        "source_branch": task.source_branch,
        "created_at": task.created_at.strftime("%d-%m-%Y %H:%M"),
    }


def _run_git(args):
    project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    try:
        return subprocess.run(
            ["git", "-C", project_root, *args],
            capture_output=True,
            text=True,
            timeout=15,
            check=False,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        return None, str(exc)


def _git_path_is_safe(path):
    normalized = str(path or "").replace("\\", "/")
    parsed = PurePosixPath(normalized)
    return bool(normalized) and not parsed.is_absolute() and ".." not in parsed.parts


@login_required(login_url="/login/")
def git_status(request):
    branch_result = _run_git(["rev-parse", "--abbrev-ref", "HEAD"])
    status_result = _run_git(["status", "--short"])
    if isinstance(branch_result, tuple) or isinstance(status_result, tuple):
        error = branch_result[1] if isinstance(branch_result, tuple) else status_result[1]
        return JsonResponse({"success": False, "error": "Git is unavailable: " + error}, status=503)
    if branch_result.returncode != 0 or status_result.returncode != 0:
        return JsonResponse({"success": False, "error": (status_result.stderr or branch_result.stderr).strip()}, status=503)

    files = []
    for line in status_result.stdout.splitlines():
        if len(line) < 3:
            continue
        path = line[3:]
        if " -> " in path:
            path = path.split(" -> ", 1)[1]
        files.append({"status": line[:2].strip() or "??", "path": path})
    return JsonResponse({
        "success": True,
        "branch": branch_result.stdout.strip() or "detached",
        "files": files,
        "raw": status_result.stdout,
    })


@login_required(login_url="/login/")
def git_diff(request):
    result = _run_git(["diff", "HEAD", "--"])
    if isinstance(result, tuple):
        return JsonResponse({"success": False, "error": "Git is unavailable: " + result[1]}, status=503)
    if result.returncode != 0:
        return JsonResponse({"success": False, "error": result.stderr.strip()}, status=503)
    return JsonResponse({"success": True, "diff": result.stdout[:MAX_GIT_OUTPUT_CHARS]})


@login_required(login_url="/login/")
@require_POST
def git_stage(request):
    paths = request.POST.getlist("paths")
    if not paths:
        paths = [request.POST.get("path", "").strip()] if request.POST.get("path") else []
    if any(not _git_path_is_safe(path) for path in paths):
        return JsonResponse({"success": False, "error": "One or more Git paths are invalid."}, status=400)
    result = _run_git(["add", "--", *paths] if paths else ["add", "-A"])
    if isinstance(result, tuple):
        return JsonResponse({"success": False, "error": "Git is unavailable: " + result[1]}, status=503)
    if result.returncode != 0:
        return JsonResponse({"success": False, "error": result.stderr.strip()}, status=400)
    return JsonResponse({"success": True, "staged": paths or ["all changes"]})


@login_required(login_url="/login/")
@require_POST
def git_commit(request):
    message = request.POST.get("message", "").strip()
    if not message:
        return JsonResponse({"success": False, "error": "Enter a commit message."}, status=400)
    if len(message) > 200:
        return JsonResponse({"success": False, "error": "Commit messages are limited to 200 characters."}, status=400)
    result = _run_git(["commit", "-m", message])
    if isinstance(result, tuple):
        return JsonResponse({"success": False, "error": "Git is unavailable: " + result[1]}, status=503)
    if result.returncode != 0:
        return JsonResponse({"success": False, "error": (result.stderr or result.stdout).strip()}, status=400)
    return JsonResponse({"success": True, "output": result.stdout.strip()})


def _remote_connection(request):
    return request.session.get("remote_git", {})


def _remote_headers(connection):
    token = connection.get("token", "")
    if connection.get("provider") == "gitlab":
        return {"PRIVATE-TOKEN": token, "Accept": "application/json"}
    return {"Authorization": "Bearer " + token, "Accept": "application/vnd.github+json"}


def _remote_request(request, method, url, **kwargs):
    preferences = PrivacyPreference.objects.filter(user=request.user).first()
    if preferences and preferences.network_lock_enabled:
        return None, JsonResponse({"success": False, "error": "Network lock is enabled. Disable it in Privacy before using remote providers."}, status=423)
    connection = _remote_connection(request)
    if not connection.get("token"):
        return None, JsonResponse({"success": False, "error": "Connect a GitHub or GitLab token first."}, status=400)
    headers = _remote_headers(connection)
    headers.update(kwargs.pop("headers", {}))
    try:
        response = requests.request(method, url, headers=headers, timeout=20, **kwargs)
    except requests.RequestException as exc:
        return None, JsonResponse({"success": False, "error": "Remote provider unavailable: " + str(exc)}, status=503)
    if not response.ok:
        return None, JsonResponse({"success": False, "error": response.text[:1000]}, status=400)
    return response, None


@login_required(login_url="/login/")
@require_http_methods(["GET", "POST"])
def remote_git_settings(request):
    connection = _remote_connection(request)
    if request.method == "POST":
        provider = request.POST.get("provider", "").strip().lower()
        if provider not in {"github", "gitlab"}:
            return JsonResponse({"success": False, "error": "Choose GitHub or GitLab."}, status=400)
        token = request.POST.get("token", "").strip()
        if token:
            connection["token"] = token
        connection["provider"] = provider
        connection["repository"] = request.POST.get("repository", "").strip()[:200]
        request.session["remote_git"] = connection
        request.session.modified = True
    return JsonResponse({
        "success": True,
        "provider": connection.get("provider", "github"),
        "repository": connection.get("repository", ""),
        "token_set": bool(connection.get("token")),
    })


@login_required(login_url="/login/")
@require_GET
def remote_repositories(request):
    connection = _remote_connection(request)
    if connection.get("provider") == "gitlab":
        url = "https://gitlab.com/api/v4/projects?membership=true&per_page=100"
    else:
        url = "https://api.github.com/user/repos?per_page=100&sort=updated"
    response, error = _remote_request(request, "GET", url)
    if error:
        return error
    items = response.json()
    repositories = [
        {
            "name": item.get("full_name") or item.get("path_with_namespace"),
            "id": item.get("id"),
            "url": item.get("html_url") or item.get("web_url"),
        }
        for item in items
    ]
    return JsonResponse({"success": True, "repositories": repositories})


@login_required(login_url="/login/")
@require_GET
def remote_issues(request):
    connection = _remote_connection(request)
    repository = request.GET.get("repository", connection.get("repository", "")).strip()
    if not repository:
        return JsonResponse({"success": False, "error": "Choose a repository first."}, status=400)
    if connection.get("provider") == "gitlab":
        url = "https://gitlab.com/api/v4/projects/" + quote(repository, safe="") + "/issues?state=opened&per_page=30"
    else:
        url = f"https://api.github.com/repos/{repository}/issues?state=open&per_page=30"
    response, error = _remote_request(request, "GET", url)
    if error:
        return error
    return JsonResponse({"success": True, "issues": [
        {"id": item.get("iid") or item.get("number"), "title": item.get("title"), "url": item.get("html_url") or item.get("web_url")}
        for item in response.json()
        if not item.get("pull_request")
    ]})


@login_required(login_url="/login/")
@require_POST
def remote_pull_request(request):
    connection = _remote_connection(request)
    repository = request.POST.get("repository", connection.get("repository", "")).strip()
    title = request.POST.get("title", "").strip()
    body = request.POST.get("body", "").strip()
    head = request.POST.get("head", "").strip()
    base = request.POST.get("base", "main").strip() or "main"
    if not repository or not title or not head:
        return JsonResponse({"success": False, "error": "Repository, title, and source branch are required."}, status=400)
    if connection.get("provider") == "gitlab":
        url = "https://gitlab.com/api/v4/projects/" + quote(repository, safe="") + "/merge_requests"
        payload = {"source_branch": head, "target_branch": base, "title": title, "description": body}
    else:
        url = f"https://api.github.com/repos/{repository}/pulls"
        payload = {"head": head, "base": base, "title": title, "body": body, "draft": True}
    response, error = _remote_request(request, "POST", url, json=payload)
    if error:
        return error
    data = response.json()
    return JsonResponse({"success": True, "url": data.get("html_url") or data.get("web_url"), "number": data.get("number") or data.get("iid")})


@csrf_exempt
@require_POST
def cli_ask(request):
    if request.META.get("REMOTE_ADDR") not in {"127.0.0.1", "::1", "localhost"}:
        return JsonResponse({"success": False, "error": "CLI access is limited to the local machine."}, status=403)
    return ask_code(request)


def _user_ollama_settings(user):
    return UserOllamaSettings.objects.get_or_create(user=user)[0]


def _safe_ollama_url(value):
    parsed = urlparse(str(value or "").strip())
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        return None
    if parsed.hostname not in ("127.0.0.1", "localhost", "::1"):
        return None
    return f"{parsed.scheme}://{parsed.netloc}".rstrip("/")


def _available_models(base_url=OLLAMA_BASE_URL):
    names = list(
        LocalModelConfig.objects.filter(is_active=True)
        .order_by("-is_default", "name")
        .values_list("name", flat=True)
    )

    if DEFAULT_MODEL not in names:
        names.insert(0, DEFAULT_MODEL)
    if VISION_MODEL not in names:
        names.append(VISION_MODEL)

    try:
        response = requests.get(f"{base_url}/api/tags", timeout=3)
        if response.ok:
            for item in response.json().get("models", []):
                name = item.get("name")
                if name and name not in names:
                    names.append(name)
    except (requests.RequestException, ValueError):
        pass

    return names


def _ollama_model_details(base_url=OLLAMA_BASE_URL):
    """Return safe local model metadata for the model manager."""
    try:
        response = requests.get(f"{base_url}/api/tags", timeout=3)
        response.raise_for_status()
        details = []
        for item in response.json().get("models", []):
            name = item.get("name")
            if not name:
                continue
            model_details = item.get("details") or {}
            details.append({
                "name": name,
                "size_bytes": int(item.get("size") or 0),
                "modified_at": item.get("modified_at") or "",
                "parameter_size": model_details.get("parameter_size", ""),
                "quantization": model_details.get("quantization_level", ""),
                "family": model_details.get("family", ""),
            })
        return details
    except (requests.RequestException, ValueError, TypeError):
        return []


def _event(payload):
    return json.dumps(payload, ensure_ascii=False) + "\n"


def _safe_filename(name):
    path = PurePosixPath(str(name).replace("\\", "/"))
    safe_parts = [part for part in path.parts if part not in ("", ".", "..")]
    return "/".join(safe_parts) or "uploaded-file"


def _language_for_filename(filename):
    extension = PurePosixPath(filename).suffix.lower().lstrip(".")
    return {
        "py": "python",
        "js": "javascript",
        "ts": "typescript",
        "jsx": "javascript",
        "tsx": "typescript",
        "html": "html",
        "css": "css",
        "sql": "sql",
        "java": "java",
        "cs": "csharp",
        "php": "php",
        "json": "json",
        "md": "markdown",
    }.get(extension, "auto")


def _extract_document_text(filename, raw_content):
    """Extract local text while preserving page markers for citations."""
    extension = PurePosixPath(filename).suffix.lower()
    if extension == ".pdf":
        if PdfReader is None:
            raise ValueError("PDF support requires the pypdf package.")
        pages = PdfReader(io.BytesIO(raw_content)).pages
        return "\n\n".join(
            f"[Page {index + 1}]\n{page.extract_text() or ''}"
            for index, page in enumerate(pages)
        ).strip()
    if extension == ".docx":
        if DocxDocument is None:
            raise ValueError("DOCX support requires the python-docx package.")
        document = DocxDocument(io.BytesIO(raw_content))
        return "\n".join(
            paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()
        ).strip()
    text = raw_content.decode("utf-8", errors="ignore")
    if "\x00" in text:
        raise ValueError("Binary files are not supported for text indexing.")
    return text


def _save_knowledge_document(filename, file_text, source_type, owner):
    content_hash = hashlib.sha256(file_text.encode("utf-8")).hexdigest()
    document = KnowledgeDocument.objects.filter(owner=owner, filename=filename).first()
    if document is None:
        document = KnowledgeDocument(owner=owner, filename=filename)
    document.title = PurePosixPath(filename).name
    document.file_extension = PurePosixPath(filename).suffix.lower()
    document.source_type = source_type
    document.language = _language_for_filename(filename)
    document.original_text = file_text
    document.file_size_bytes = len(file_text.encode("utf-8"))
    document.content_hash = content_hash
    document.is_active = True
    document.save()
    _index_knowledge_document(document, owner)
    return document


def _ollama_embeddings(texts, base_url):
    if not texts:
        return []
    embedding_model = os.environ.get("OLLAMA_EMBED_MODEL", "nomic-embed-text")
    try:
        response = requests.post(
            f"{base_url}/api/embed",
            json={"model": embedding_model, "input": texts},
            timeout=(5, 60),
        )
        response.raise_for_status()
        payload = response.json()
        embeddings = payload.get("embeddings") or []
        if not embeddings and payload.get("embedding"):
            embeddings = [payload["embedding"]]
        return embeddings if len(embeddings) == len(texts) else []
    except (requests.RequestException, ValueError, TypeError):
        return []


def _index_knowledge_document(document, owner):
    document.chunks.all().delete()
    indexed_text = redact_sensitive_text(owner, document.original_text)
    contents = [
        indexed_text[start:start + CHUNK_SIZE]
        for start in range(0, len(indexed_text), CHUNK_SIZE)
    ]
    settings = _user_ollama_settings(owner) if owner else None
    embeddings = _ollama_embeddings(contents, settings.server_url if settings else OLLAMA_BASE_URL)
    KnowledgeChunk.objects.bulk_create([
        KnowledgeChunk(
            document=document,
            chunk_index=index,
            content=content,
            language=document.language,
            embedding=embeddings[index] if index < len(embeddings) else [],
        )
        for index, content in enumerate(contents)
    ])


def _cosine_similarity(left, right):
    if not left or not right or len(left) != len(right):
        return 0.0
    numerator = sum(a * b for a, b in zip(left, right))
    left_norm = sum(a * a for a in left) ** 0.5
    right_norm = sum(b * b for b in right) ** 0.5
    return numerator / (left_norm * right_norm) if left_norm and right_norm else 0.0


def _search_knowledge(query, limit=8, owner=None, filenames=None):
    terms = list(dict.fromkeys(re.findall(r"[a-zA-Z0-9_]{2,}", query.lower())))
    selected_filenames = {str(name).strip().lower() for name in (filenames or []) if str(name).strip()}
    if not terms and not selected_filenames:
        return []

    user_settings = _user_ollama_settings(owner) if owner else None
    query_embedding = _ollama_embeddings([query], user_settings.server_url if user_settings else OLLAMA_BASE_URL)
    query_embedding = query_embedding[0] if query_embedding else []
    matches = []
    chunks = KnowledgeChunk.objects.filter(
        document__is_active=True,
        document__owner=owner,
    ).select_related("document")

    for chunk in chunks:
        content = chunk.content.lower()
        filename = (chunk.document.filename or "").lower()
        if selected_filenames and filename not in selected_filenames:
            continue
        lexical_score = sum(content.count(term) for term in terms)
        semantic_score = _cosine_similarity(query_embedding, chunk.embedding)
        score = lexical_score + (semantic_score * 100)
        score += 3 * sum(filename.count(term) for term in terms)
        if score or selected_filenames:
            matches.append((max(score, 1), chunk))

        matches.sort(key=lambda item: item[0], reverse=True)
    return [
        {
            "filename": chunk.document.filename,
            "language": chunk.language,
            "chunk_index": chunk.chunk_index,
            "content": chunk.content,
            "score": score,
            "line_start": (
                chunk.document.original_text[:max(
                    chunk.document.original_text.find(chunk.content),
                    0,
                )].count("\n") + 1
            ),
            "line_end": (
                chunk.document.original_text[:max(
                    chunk.document.original_text.find(chunk.content),
                    0,
                )].count("\n") + chunk.content.count("\n") + 1
            ),
            "page_start": _page_for_offset(chunk.document.original_text, chunk.document.original_text.find(chunk.content)),
            "page_end": _page_for_offset(chunk.document.original_text, chunk.document.original_text.find(chunk.content) + len(chunk.content)),
        }
        for score, chunk in matches[:limit]
    ]


def _page_for_offset(text, offset):
    if offset < 0:
        return None
    pages = re.findall(r"\[Page\s+(\d+)\]", text[:offset])
    return int(pages[-1]) if pages else None


def signup(request):
    if request.user.is_authenticated:
        return redirect("chat_home")
    if request.method == "POST":
        username = request.POST.get("username", "").strip()
        password = request.POST.get("password", "")
        confirmation = request.POST.get("confirmation", "")
        if not username or not password:
            return render(request, "chat/signup.html", {"error": "Username and password are required."})
        if password != confirmation:
            return render(request, "chat/signup.html", {"error": "Passwords do not match."})
        if User.objects.filter(username=username).exists():
            return render(request, "chat/signup.html", {"error": "That username is already in use."})
        user = User.objects.create_user(username=username, password=password)
        login(request, user)
        return redirect("chat_home")
    return render(request, "chat/signup.html")


def login_view(request):
    if request.user.is_authenticated:
        return redirect("chat_home")
    if request.method == "POST":
        user = authenticate(
            request,
            username=request.POST.get("username", "").strip(),
            password=request.POST.get("password", ""),
        )
        if user is not None:
            login(request, user)
            return redirect("chat_home")
        return render(request, "chat/login.html", {"error": "Invalid username or password."})
    return render(request, "chat/login.html")


def _sso_config(provider, workspace_id=None):
    query = EnterpriseIdentityConfig.objects.filter(provider=provider)
    if workspace_id:
        try:
            query = query.filter(workspace_id=int(workspace_id))
        except (TypeError, ValueError):
            return None
    return query.select_related("workspace").order_by("workspace_id").first()


def _sso_domain_allowed(config, email):
    domains = [item.strip().lower().lstrip("@") for item in config.allowed_domains.split(",") if item.strip()]
    if not domains:
        return True
    email = email.strip().lower()
    return any(email.endswith("@" + domain) for domain in domains)


def _sso_login_user(request, config, claims):
    email = str(claims.get("email") or claims.get("upn") or "").strip().lower()
    if not email or "@" not in email:
        raise ValueError("The identity provider did not return an email address.")
    if not _sso_domain_allowed(config, email):
        raise ValueError("This email domain is not allowed for the workspace.")
    username = str(claims.get("preferred_username") or email.split("@", 1)[0]).strip()
    username = re.sub(r"[^A-Za-z0-9_.@+-]", "-", username)[:150] or email.split("@", 1)[0][:150]
    user, created = User.objects.get_or_create(username=username, defaults={"email": email})
    changed = []
    if user.email != email:
        user.email = email
        changed.append("email")
    display_name = str(claims.get("name") or claims.get("display_name") or "").strip()
    if display_name:
        parts = display_name.split(" ", 1)
        if user.first_name != parts[0][:150]:
            user.first_name = parts[0][:150]
            changed.append("first_name")
        last_name = parts[1][:150] if len(parts) > 1 else ""
        if user.last_name != last_name:
            user.last_name = last_name
            changed.append("last_name")
    if not user.is_active:
        user.is_active = True
        changed.append("is_active")
    if created:
        user.set_unusable_password()
        user.save()
    elif changed:
        user.save(update_fields=changed)
    WorkspaceMembership.objects.get_or_create(workspace=config.workspace, user=user, defaults={"role": "developer"})
    login(request, user, backend="django.contrib.auth.backends.ModelBackend")
    return user


@require_GET
def sso_login(request, provider):
    provider = provider.strip().lower()
    if provider not in {"oidc", "saml"}:
        return JsonResponse({"success": False, "error": "Unsupported SSO provider."}, status=400)
    config = _sso_config(provider, request.GET.get("workspace_id"))
    if not config:
        return JsonResponse({"success": False, "error": "No configured workspace identity provider found."}, status=404)
    state = secrets.token_urlsafe(32)
    callback = request.build_absolute_uri(reverse("sso_callback", kwargs={"provider": provider}))
    request.session[f"sso:{state}"] = {"workspace_id": config.workspace_id, "callback": callback}
    request.session.set_expiry(600)
    if provider == "saml":
        if not config.saml_entrypoint_url:
            return JsonResponse({"success": False, "error": "SAML entrypoint is not configured."}, status=400)
        suffix = "&" if "?" in config.saml_entrypoint_url else "?"
        return redirect(config.saml_entrypoint_url + suffix + urlencode({"RelayState": state, "redirect_uri": callback}))
    if not config.issuer_url or not config.client_id:
        return JsonResponse({"success": False, "error": "OIDC issuer URL and client ID are required."}, status=400)
    try:
        discovery = requests.get(config.issuer_url.rstrip("/") + "/.well-known/openid-configuration", timeout=10)
        discovery.raise_for_status()
        metadata = discovery.json()
        authorization_endpoint = metadata["authorization_endpoint"]
    except (requests.RequestException, KeyError, ValueError) as exc:
        return JsonResponse({"success": False, "error": f"OIDC discovery failed: {exc}"}, status=502)
    verifier = secrets.token_urlsafe(48)
    nonce = secrets.token_urlsafe(24)
    request.session[f"sso:{state}"].update({"verifier": verifier, "nonce": nonce})
    challenge = base64.urlsafe_b64encode(hashlib.sha256(verifier.encode("utf-8")).digest()).decode("ascii").rstrip("=")
    params = {
        "client_id": config.client_id,
        "response_type": "code",
        "redirect_uri": callback,
        "scope": "openid profile email",
        "state": state,
        "nonce": nonce,
        "code_challenge": challenge,
        "code_challenge_method": "S256",
    }
    return redirect(authorization_endpoint + "?" + urlencode(params))


@require_http_methods(["GET", "POST"])
def sso_callback(request, provider):
    provider = provider.strip().lower()
    state = request.GET.get("state") or request.POST.get("RelayState")
    state_data = request.session.pop(f"sso:{state}", None) if state else None
    if not state_data:
        return JsonResponse({"success": False, "error": "SSO state is missing or expired."}, status=400)
    config = _sso_config(provider, state_data.get("workspace_id"))
    if not config:
        return JsonResponse({"success": False, "error": "SSO workspace configuration was not found."}, status=404)
    if provider == "saml":
        return JsonResponse({"success": False, "error": "SAML assertion validation must be handled by a signed SAML gateway before this callback."}, status=501)
    if request.GET.get("error"):
        return JsonResponse({"success": False, "error": request.GET.get("error_description") or request.GET.get("error")}, status=400)
    code = request.GET.get("code", "").strip()
    if not code:
        return JsonResponse({"success": False, "error": "OIDC authorization code is missing."}, status=400)
    callback = state_data.get("callback") or request.build_absolute_uri(reverse("sso_callback", kwargs={"provider": provider}))
    try:
        discovery = requests.get(config.issuer_url.rstrip("/") + "/.well-known/openid-configuration", timeout=10)
        discovery.raise_for_status()
        metadata = discovery.json()
        token_response = requests.post(metadata["token_endpoint"], data={
            "grant_type": "authorization_code",
            "code": code,
            "client_id": config.client_id,
            "redirect_uri": callback,
            "code_verifier": state_data.get("verifier", ""),
        }, timeout=10)
        token_response.raise_for_status()
        token_data = token_response.json()
        access_token = token_data.get("access_token")
        if not access_token or not metadata.get("userinfo_endpoint"):
            raise ValueError("OIDC provider did not return an access token and userinfo endpoint.")
        userinfo = requests.get(metadata["userinfo_endpoint"], headers={"Authorization": "Bearer " + access_token}, timeout=10)
        userinfo.raise_for_status()
        _sso_login_user(request, config, userinfo.json())
    except (requests.RequestException, KeyError, ValueError) as exc:
        return JsonResponse({"success": False, "error": f"OIDC sign-in failed: {exc}"}, status=502)
    return redirect("chat_home")


@require_POST
def logout_view(request):
    logout(request)
    return redirect("chat_home")


def index(request):
    sessions = (
        ChatSession.objects.filter(owner=request.user, is_archived=False)
        .order_by("-is_pinned", "-updated_at")[:30]
        if request.user.is_authenticated
        else ChatSession.objects.none()
    )
    user_settings = _user_ollama_settings(request.user) if request.user.is_authenticated else None
    server_url = user_settings.server_url if user_settings else OLLAMA_BASE_URL
    default_model = user_settings.default_model if user_settings else DEFAULT_MODEL
    return render(request, "chat/index.html", {
        "sessions": sessions,
        "models": _available_models(server_url),
        "default_model": default_model,
        "ollama_settings": user_settings,
    })


def model_list(request):
    user_settings = _user_ollama_settings(request.user) if request.user.is_authenticated else None
    server_url = user_settings.server_url if user_settings else OLLAMA_BASE_URL
    default_model = user_settings.default_model if user_settings else DEFAULT_MODEL
    return JsonResponse({
        "success": True,
        "models": _available_models(server_url),
        "model_details": _ollama_model_details(server_url),
        "default_model": default_model,
    })


def _ollama_settings_payload(settings):
    return {
        "server_url": settings.server_url,
        "default_model": settings.default_model,
        "fallback_model": settings.fallback_model,
        "temperature": settings.temperature,
        "top_p": settings.top_p,
        "max_context_chars": settings.max_context_chars,
    }


@login_required(login_url="/login/")
@require_http_methods(["GET", "POST"])
def ollama_settings(request):
    settings = _user_ollama_settings(request.user)
    if request.method == "POST":
        server_url = _safe_ollama_url(request.POST.get("server_url", settings.server_url))
        if not server_url:
            return JsonResponse({"success": False, "error": "Use a local Ollama URL such as http://127.0.0.1:11434."}, status=400)
        try:
            temperature = min(max(float(request.POST.get("temperature", settings.temperature)), 0), 2)
            top_p = min(max(float(request.POST.get("top_p", settings.top_p)), 0), 1)
            max_context_chars = min(max(int(request.POST.get("max_context_chars", settings.max_context_chars)), 4000), 100000)
        except (TypeError, ValueError):
            return JsonResponse({"success": False, "error": "Temperature, top-p, and context must be valid numbers."}, status=400)
        settings.server_url = server_url
        settings.default_model = request.POST.get("default_model", settings.default_model).strip()[:100] or DEFAULT_MODEL
        settings.fallback_model = request.POST.get("fallback_model", settings.fallback_model).strip()[:100]
        settings.temperature = temperature
        settings.top_p = top_p
        settings.max_context_chars = max_context_chars
        settings.save()
    return JsonResponse({
        "success": True,
        "settings": _ollama_settings_payload(settings),
        "models": _available_models(settings.server_url),
        "model_details": _ollama_model_details(settings.server_url),
    })


@login_required(login_url="/login/")
@require_GET
def ollama_health(request):
    settings = _user_ollama_settings(request.user)
    started = time.perf_counter()
    try:
        response = requests.get(f"{settings.server_url}/api/tags", timeout=5)
        response.raise_for_status()
        models = [
            item.get("name")
            for item in response.json().get("models", [])
            if item.get("name")
        ]
        return JsonResponse({
            "success": True,
            "status": "ready",
            "server_url": settings.server_url,
            "latency_ms": round((time.perf_counter() - started) * 1000),
            "models": models,
            "model_details": _ollama_model_details(settings.server_url),
            "default_model": settings.default_model,
            "fallback_model": settings.fallback_model,
            "fallback_available": settings.fallback_model in models,
        })
    except (requests.RequestException, ValueError) as exc:
        return JsonResponse({
            "success": False,
            "status": "offline",
            "server_url": settings.server_url,
            "latency_ms": round((time.perf_counter() - started) * 1000),
            "error": "Ollama health check failed: " + str(exc),
        }, status=503)

@login_required(login_url="/login/")
@require_POST
def ollama_model_action(request):
    settings = _user_ollama_settings(request.user)
    action = request.POST.get("action", "").strip().lower()
    name = request.POST.get("name", "").strip()
    if not name or len(name) > 100:
        return JsonResponse({"success": False, "error": "Enter a valid model name."}, status=400)
    try:
        if action == "pull":
            response = requests.post(
                f"{settings.server_url}/api/pull",
                json={"name": name, "stream": True},
                timeout=(10, 300),
            )
        elif action == "delete":
            response = requests.delete(
                f"{settings.server_url}/api/delete",
                json={"name": name},
                timeout=(10, 30),
            )
        else:
            return JsonResponse({"success": False, "error": "Use pull or delete."}, status=400)
    except requests.RequestException as exc:
        return JsonResponse({"success": False, "error": "Ollama is unavailable: " + str(exc)}, status=503)
    if not response.ok:
        return JsonResponse({"success": False, "error": response.text}, status=400)

    progress = []
    if action == "pull" and hasattr(response, "iter_lines"):
        for raw_line in response.iter_lines(decode_unicode=True):
            if not raw_line:
                continue
            try:
                event = json.loads(raw_line)
            except (TypeError, ValueError):
                continue
            progress.append({
                "status": event.get("status", ""),
                "completed": event.get("completed", 0),
                "total": event.get("total", 0),
            })
    return JsonResponse({
        "success": True,
        "models": _available_models(settings.server_url),
        "model_details": _ollama_model_details(settings.server_url),
        "progress": progress[-100:],
    })


@login_required(login_url="/login/")
def session_list(request):
    query = request.GET.get("q", "").strip()
    sessions = ChatSession.objects.filter(
        owner=request.user,
        is_archived=False,
    )
    if query:
        sessions = sessions.filter(
            Q(title__icontains=query) | Q(messages__content__icontains=query)
        ).distinct()
    sessions = sessions.order_by("-is_pinned", "-updated_at")[:100]
    return JsonResponse({
        "success": True,
        "sessions": [
            {
                "id": session.id,
                "title": session.title,
                "is_pinned": session.is_pinned,
                "tags": [tag.strip() for tag in session.tags.split(",") if tag.strip()],
                "updated_at": session.updated_at.strftime("%d-%m-%Y %H:%M"),
            }
            for session in sessions
        ],
    })


@login_required(login_url="/login/")
@require_POST
def manage_session(request, session_id):
    session = get_object_or_404(ChatSession, id=session_id, owner=request.user)
    action = request.POST.get("action", "").strip().lower()

    if action == "rename":
        title = request.POST.get("title", "").strip()[:200]
        if not title:
            return JsonResponse({"success": False, "error": "A chat title is required."}, status=400)
        session.title = title
    elif action == "tag":
        tags = request.POST.get("tags", "")
        session.tags = ",".join(dict.fromkeys(
            tag.strip()[:40] for tag in tags.split(",") if tag.strip()
        ))[:300]
    elif action == "pin":
        session.is_pinned = not session.is_pinned
    elif action == "archive":
        session.is_archived = True
    elif action == "clear_context":
        session.messages.all().delete()
        session.save(update_fields=["updated_at"])
        return JsonResponse({"success": True, "id": session.id, "cleared": True})
    elif action == "delete":
        session.delete()
        return JsonResponse({"success": True, "deleted": True})
    else:
        return JsonResponse({"success": False, "error": "Unsupported chat action."}, status=400)

    session.save()
    return JsonResponse({
        "success": True,
        "id": session.id,
        "title": session.title,
        "is_pinned": session.is_pinned,
        "tags": [tag.strip() for tag in session.tags.split(",") if tag.strip()],
    })


@login_required(login_url="/login/")
def knowledge_search(request):
    query = request.GET.get("q", "").strip()
    return JsonResponse({
        "success": True,
        "query": query,
        "results": _search_knowledge(query, owner=request.user),
    })


@login_required(login_url="/login/")
def project_documents(request):
    documents = (
        KnowledgeDocument.objects.filter(owner=request.user, is_active=True)
        .order_by("filename")
    )
    return JsonResponse({
        "success": True,
        "documents": [
            {
                "id": document.id,
                "filename": document.filename or document.title,
                "language": document.language,
                "size_bytes": document.file_size_bytes,
                "chunks": document.chunks.count(),
                "uploaded_at": document.uploaded_at.strftime("%d-%m-%Y %H:%M"),
            }
            for document in documents
        ],
    })


@login_required(login_url="/login/")
@require_POST
def project_document_create(request):
    filename = _safe_filename(request.POST.get("filename", "generated-file"))
    content = request.POST.get("content", "")
    if not content.strip():
        return JsonResponse({"success": False, "error": "Generated content is empty."}, status=400)
    if len(content.encode("utf-8")) > MAX_FILE_BYTES:
        return JsonResponse({"success": False, "error": "Saved files are limited to 1 MB."}, status=400)
    document = _save_knowledge_document(filename, content, "manual", request.user)
    return JsonResponse({
        "success": True,
        "document": {
            "id": document.id,
            "filename": document.filename,
            "language": document.language,
            "size_bytes": document.file_size_bytes,
        },
    })


@login_required(login_url="/login/")
@require_POST
def local_ocr(request):
    uploaded_files = request.FILES.getlist("files")
    if not uploaded_files:
        return JsonResponse({"success": False, "error": "Choose an image or scanned PDF first."}, status=400)
    settings = _user_ollama_settings(request.user)
    image_data = []
    names = []
    for uploaded_file in uploaded_files[:5]:
        if uploaded_file.size > MAX_IMAGE_BYTES * 2:
            return JsonResponse({"success": False, "error": f"{uploaded_file.name} is too large for local OCR."}, status=400)
        raw = uploaded_file.read()
        extension = PurePosixPath(uploaded_file.name).suffix.lower()
        if extension == ".pdf":
            if fitz is None:
                return JsonResponse({"success": False, "error": "Scanned PDF OCR requires the PyMuPDF package."}, status=400)
            document = fitz.open(stream=raw, filetype="pdf")
            for page_index in range(min(5, document.page_count)):
                page = document.load_page(page_index)
                pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                image_data.append(base64.b64encode(pixmap.tobytes("png")).decode("ascii"))
            names.append(uploaded_file.name)
        elif uploaded_file.content_type.startswith("image/"):
            image_data.append(base64.b64encode(raw).decode("ascii"))
            names.append(uploaded_file.name)
        else:
            return JsonResponse({"success": False, "error": f"{uploaded_file.name} is not an image or PDF."}, status=400)
    try:
        response = requests.post(
            f"{settings.server_url}/api/generate",
            json={
                "model": VISION_MODEL,
                "prompt": "Transcribe all visible text exactly. Preserve headings, line breaks, tables, and code. Return only the transcription.",
                "images": image_data,
                "stream": False,
            },
            timeout=(10, 300),
        )
        response.raise_for_status()
        text = response.json().get("response", "").strip()
        return JsonResponse({"success": True, "files": names, "text": text, "pages": len(image_data)})
    except (requests.RequestException, ValueError, TypeError) as exc:
        return JsonResponse({"success": False, "error": "Local OCR failed: " + str(exc)}, status=503)


@login_required(login_url="/login/")
@require_POST
def local_transcribe(request):
    global _whisper_model
    audio = request.FILES.get("audio")
    if not audio:
        return JsonResponse({"success": False, "error": "Record or choose an audio file first."}, status=400)
    if WhisperModel is None:
        return JsonResponse({"success": False, "error": "Offline speech-to-text requires the faster-whisper package."}, status=400)
    if audio.size > 25 * 1024 * 1024:
        return JsonResponse({"success": False, "error": "Audio is limited to 25 MB."}, status=400)
    try:
        if _whisper_model is None:
            _whisper_model = WhisperModel(
                os.environ.get("WHISPER_MODEL", "base"),
                device=os.environ.get("WHISPER_DEVICE", "cpu"),
                compute_type=os.environ.get("WHISPER_COMPUTE_TYPE", "int8"),
            )
        with tempfile.NamedTemporaryFile(suffix=PurePosixPath(audio.name).suffix or ".webm") as temporary:
            for chunk in audio.chunks():
                temporary.write(chunk)
            temporary.flush()
            segments, info = _whisper_model.transcribe(temporary.name, vad_filter=True)
            text = " ".join(segment.text.strip() for segment in segments).strip()
        return JsonResponse({"success": True, "text": text, "language": getattr(info, "language", "")})
    except Exception as exc:
        return JsonResponse({"success": False, "error": "Local transcription failed: " + str(exc)}, status=503)


@login_required(login_url="/login/")
@require_http_methods(["GET", "POST"])
def privacy_dashboard(request):
    preferences, _ = PrivacyPreference.objects.get_or_create(user=request.user)
    if request.method == "POST":
        preferences.network_lock_enabled = request.POST.get("network_lock_enabled", "true").lower() == "true"
        preferences.store_chat_history = request.POST.get("store_chat_history", "true").lower() == "true"
        preferences.redact_secrets = request.POST.get("redact_secrets", "true").lower() == "true"
        preferences.save()
    return JsonResponse({
        "success": True,
        "privacy": {
            "network_lock_enabled": preferences.network_lock_enabled,
            "store_chat_history": preferences.store_chat_history,
            "redact_secrets": preferences.redact_secrets,
            "updated_at": preferences.updated_at.isoformat(),
        },
        "runtime": {
            "provider": "Ollama on localhost",
            "database_path": str(settings.DATABASES["default"].get("NAME", "")),
            "media_path": str(getattr(settings, "MEDIA_ROOT", "")),
            "remote_token_configured": bool(_remote_connection(request).get("token")),
            "network_lock_effect": "Remote provider calls are blocked" if preferences.network_lock_enabled else "Remote provider calls require explicit connector settings",
        },
    })

@login_required(login_url="/login/")
@require_http_methods(["DELETE"])
def project_document_delete(request, document_id):
    document = get_object_or_404(
        KnowledgeDocument,
        id=document_id,
        owner=request.user,
    )
    document.delete()
    return JsonResponse({"success": True, "deleted_id": document_id})


@login_required(login_url="/login/")
@require_http_methods(["GET", "POST"])
def project_document_content(request, document_id):
    document = get_object_or_404(
        KnowledgeDocument,
        id=document_id,
        owner=request.user,
    )
    if request.method == "GET":
        return JsonResponse({
            "success": True,
            "document": {
                "id": document.id,
                "filename": document.filename or document.title,
                "language": document.language,
                "content": document.original_text,
            },
        })

    content = request.POST.get("content", "")
    if len(content.encode("utf-8")) > MAX_FILE_BYTES:
        return JsonResponse({"success": False, "error": "Saved files are limited to 1 MB."}, status=400)
    document.original_text = content
    document.file_size_bytes = len(content.encode("utf-8"))
    document.content_hash = hashlib.sha256(content.encode("utf-8")).hexdigest()
    document.save(update_fields=["original_text", "file_size_bytes", "content_hash"])
    _index_knowledge_document(document, request.user)
    return JsonResponse({
        "success": True,
        "document_id": document.id,
        "chunks": document.chunks.count(),
        "saved_at": document.uploaded_at.strftime("%d-%m-%Y %H:%M"),
    })


@login_required(login_url="/login/")
@require_POST
def project_document_reindex(request, document_id):
    document = get_object_or_404(
        KnowledgeDocument,
        id=document_id,
        owner=request.user,
    )
    document.chunks.all().delete()
    KnowledgeChunk.objects.bulk_create([
        KnowledgeChunk(
            document=document,
            chunk_index=index,
            content=document.original_text[start:start + CHUNK_SIZE],
            language=document.language,
        )
        for index, start in enumerate(range(0, len(document.original_text), CHUNK_SIZE))
    ])
    return JsonResponse({
        "success": True,
        "document_id": document.id,
        "chunks": document.chunks.count(),
    })


def session_messages(request, session_id):
    owner = request.user if request.user.is_authenticated else None
    session = get_object_or_404(ChatSession, id=session_id, owner=owner)
    messages = [
        {
            "id": msg.id,
            "role": msg.role,
            "content": msg.content,
            "filename": msg.filename,
            "created_at": msg.created_at.strftime("%d-%m-%Y %H:%M"),
        }
        for msg in session.messages.order_by("created_at")
    ]
    return JsonResponse({
        "success": True,
        "session_id": session.id,
        "title": session.title,
        "model": session.model_name,
        "messages": messages,
    })


@login_required(login_url="/login/")
@require_GET
def session_summary(request, session_id):
    session = get_object_or_404(ChatSession, id=session_id, owner=request.user)
    messages = list(session.messages.order_by("created_at"))
    user_requests = []
    for message in messages:
        if message.role != "user":
            continue
        cleaned = message.content.strip()
        for marker in ("Prompt:", "Code:"):
            if cleaned.startswith(marker):
                cleaned = cleaned[len(marker):].strip()
        user_requests.append(cleaned.replace(chr(10), " "))
    assistant_count = sum(1 for message in messages if message.role == "assistant")
    if user_requests:
        highlights = chr(10).join(f"- {request[:240]}" for request in user_requests[-8:])
        summary = (
            f"Conversation memory: {len(messages)} messages, {assistant_count} assistant responses."
            + chr(10) + "Recent requests:" + chr(10) + highlights
        )
    else:
        summary = "Conversation memory is empty. Send a message to build context."
    return JsonResponse({
        "success": True,
        "session_id": session.id,
        "summary": summary,
        "message_count": len(messages),
    })


@login_required(login_url="/login/")
def session_revisions(request, session_id):
    session = get_object_or_404(ChatSession, id=session_id, owner=request.user)
    revisions = [
        {
            "branch_id": str(branch_id),
            "created_at": created_at.strftime("%d-%m-%Y %H:%M"),
            "messages": count,
        }
        for branch_id, created_at, count in (
            ConversationRevision.objects.filter(session=session)
            .values("branch_id")
            .annotate(
                created_at=Min("created_at"),
                count=Count("id"),
            )
            .values_list("branch_id", "created_at", "count")
            .order_by("-created_at")
        )
    ]
    return JsonResponse({"success": True, "revisions": revisions})


def _session_markdown(session):
    lines = [
        f"# {session.title}",
        "",
        f"- Model: `{session.model_name}`",
        f"- Created: {session.created_at:%Y-%m-%d %H:%M}",
        "",
    ]
    for message in session.messages.order_by("created_at"):
        label = "You" if message.role == "user" else "Syntax Local AI"
        lines.extend([f"## {label}", "", message.content, ""])
    return "\n".join(lines)


def _pdf_escape(value):
    return value.replace("\\", "\\\\").replace("(", "\\\\(").replace(")", "\\\\)")


def _session_pdf(session):
    text = _session_markdown(session).encode("latin-1", errors="replace").decode("latin-1")
    lines = []
    for paragraph in text.splitlines():
        while len(paragraph) > 95:
            lines.append(paragraph[:95])
            paragraph = paragraph[95:]
        lines.append(paragraph)

    content_lines = ["BT", "/F1 10 Tf", "50 760 Td"]
    for line in lines[:65]:
        content_lines.append(f"({_pdf_escape(line)}) Tj")
        content_lines.append("0 -11 Td")
    content_lines.append("ET")
    content = "\n".join(content_lines).encode("latin-1", errors="replace")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
    ]
    pdf = b"%PDF-1.4\n%\\xe2\\xe3\\xcf\\xd3\n"
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_offset = len(pdf)
    pdf += f"xref\n0 {len(objects) + 1}\n".encode()
    pdf += b"0000000000 65535 f \n"
    pdf += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    pdf += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode()
    return pdf


@login_required(login_url="/login/")
@require_http_methods(["GET", "POST"])
def project_backup(request):
    if request.method == "GET":
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as bundle:
            bundle.writestr(
                "manifest.json",
                json.dumps({
                    "format": "syntax-local-ai-backup",
                    "version": 1,
                    "username": request.user.username,
                    "created_at": timezone.now().isoformat(),
                }, indent=2),
            )
            for document in KnowledgeDocument.objects.filter(
                owner=request.user, is_active=True
            ).order_by("filename"):
                filename = _safe_filename(document.filename or document.title)
                bundle.writestr("project/" + filename, document.original_text)
            for session in ChatSession.objects.filter(
                owner=request.user
            ).order_by("id"):
                messages = list(session.messages.order_by("created_at").values(
                    "role", "content", "filename", "model_name", "created_at"
                ))
                bundle.writestr(
                    f"chats/{session.id}.json",
                    json.dumps({
                        "title": session.title,
                        "model_name": session.model_name,
                        "tags": session.tags,
                        "messages": messages,
                    }, default=str, indent=2),
                )
            settings = _user_ollama_settings(request.user)
            bundle.writestr(
                "settings/ollama.json",
                json.dumps({
                    "server_url": settings.server_url,
                    "default_model": settings.default_model,
                    "fallback_model": settings.fallback_model,
                    "temperature": settings.temperature,
                    "top_p": settings.top_p,
                    "max_context_chars": settings.max_context_chars,
                }, indent=2),
            )
        archive_bytes = archive.getvalue()
        encrypted = request.GET.get("encrypted", "").lower() in {"1", "true", "yes"}
        response = HttpResponse(encrypt_blob(archive_bytes) if encrypted else archive_bytes, content_type="application/octet-stream" if encrypted else "application/zip")
        response["Content-Disposition"] = 'attachment; filename="syntax-local-ai-backup.enc"' if encrypted else 'attachment; filename="syntax-local-ai-backup.zip"'
        return response

    upload = request.FILES.get("backup")
    if not upload:
        return JsonResponse({"success": False, "error": "Choose a backup ZIP file."}, status=400)
    if upload.size > MAX_BACKUP_BYTES:
        return JsonResponse({"success": False, "error": "Backup files are limited to 25 MB."}, status=400)

    try:
        with zipfile.ZipFile(io.BytesIO(decrypt_blob(upload.read()))) as bundle:
            names = bundle.namelist()
            for name in names:
                parts = PurePosixPath(name).parts
                if name.startswith("/") or ".." in parts:
                    return JsonResponse({"success": False, "error": "Backup contains an unsafe path."}, status=400)

            project_count = 0
            chat_count = 0
            for name in names:
                if name.startswith("project/") and not name.endswith("/"):
                    filename = _safe_filename(name[len("project/"):])
                    content = bundle.read(name).decode("utf-8", errors="ignore")
                    if content and project_count < MAX_PROJECT_FILES:
                        _save_knowledge_document(filename, content, "manual", request.user)
                        project_count += 1

            settings_name = "settings/ollama.json"
            if settings_name in names:
                payload = json.loads(bundle.read(settings_name).decode("utf-8"))
                settings = _user_ollama_settings(request.user)
                server_url = _safe_ollama_url(payload.get("server_url", settings.server_url))
                if server_url:
                    settings.server_url = server_url
                settings.default_model = str(payload.get("default_model", settings.default_model))[:100] or DEFAULT_MODEL
                settings.fallback_model = str(payload.get("fallback_model", settings.fallback_model))[:100]
                settings.temperature = min(max(float(payload.get("temperature", settings.temperature)), 0), 2)
                settings.top_p = min(max(float(payload.get("top_p", settings.top_p)), 0), 1)
                settings.max_context_chars = min(max(int(payload.get("max_context_chars", settings.max_context_chars)), 4000), 100000)
                settings.save()

            for name in names:
                if not name.startswith("chats/") or not name.endswith(".json"):
                    continue
                payload = json.loads(bundle.read(name).decode("utf-8"))
                session = ChatSession.objects.create(
                    owner=request.user,
                    title=str(payload.get("title", "Restored chat"))[:200] or "Restored chat",
                    model_name=str(payload.get("model_name", DEFAULT_MODEL))[:100] or DEFAULT_MODEL,
                    tags=str(payload.get("tags", ""))[:300],
                )
                restored_messages = []
                for message in payload.get("messages", []):
                    role = message.get("role")
                    if role not in {"user", "assistant", "system"}:
                        continue
                    restored_messages.append(ChatMessage(
                        session=session,
                        role=role,
                        content=str(message.get("content", "")),
                        filename=str(message.get("filename", ""))[:500],
                        model_name=str(message.get("model_name", ""))[:100] or None,
                    ))
                ChatMessage.objects.bulk_create(restored_messages)
                chat_count += 1

            return JsonResponse({
                "success": True,
                "project_files": project_count,
                "chats": chat_count,
            })
    except (zipfile.BadZipFile, UnicodeDecodeError, ValueError, TypeError, OverflowError) as exc:
        return JsonResponse({"success": False, "error": "Unable to restore backup: " + str(exc)}, status=400)

@login_required(login_url="/login/")
def export_session(request, session_id):
    session = get_object_or_404(ChatSession, id=session_id, owner=request.user)
    export_format = request.GET.get("format", "markdown").lower()
    filename = f"syntax-local-ai-{session.id}"

    if export_format in ("md", "markdown"):
        response = HttpResponse(_session_markdown(session), content_type="text/markdown")
        response["Content-Disposition"] = f'attachment; filename="{filename}.md"'
        return response

    if export_format == "json":
        payload = {
            "session_id": session.id,
            "title": session.title,
            "model": session.model_name,
            "messages": list(session.messages.order_by("created_at").values(
                "role", "content", "filename", "model_name", "created_at"
            )),
        }
        response = HttpResponse(
            json.dumps(payload, default=str, indent=2),
            content_type="application/json",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}.json"'
        return response

    if export_format == "pdf":
        response = HttpResponse(_session_pdf(session), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename}.pdf"'
        return response

    if export_format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["role", "content", "filename", "model", "created_at"])
        for message in session.messages.order_by("created_at"):
            writer.writerow([message.role, message.content, message.filename or "", message.model_name or "", message.created_at.isoformat()])
        response = HttpResponse(output.getvalue(), content_type="text/csv")
        response["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
        return response

    if export_format == "code":
        code = "\n\n".join(message.content for message in session.messages.filter(role="assistant").order_by("created_at"))
        response = HttpResponse(code, content_type="text/plain")
        response["Content-Disposition"] = f'attachment; filename="{filename}-answers.txt"'
        return response

    return JsonResponse({"success": False, "error": "Unsupported export format."}, status=400)


def _javascript_check(code):
    sanitized = re.sub(r"(['\"]).*?\\1", "", code, flags=re.DOTALL)
    stack = []
    pairs = {")": "(", "]": "[", "}": "{",
    }
    for character in sanitized:
        if character in "([{":
            stack.append(character)
        elif character in pairs:
            if not stack or stack.pop() != pairs[character]:
                return False, "Unbalanced JavaScript brackets."
    return (not stack), "Unbalanced JavaScript brackets." if stack else ""


@login_required(login_url="/login/")
@require_POST
def execute_code(request):
    language = request.POST.get("language", "python").strip().lower()
    code = request.POST.get("code", "")
    if not code.strip():
        return JsonResponse({"success": False, "error": "Enter code to check."}, status=400)
    if len(code) > MAX_EXECUTION_CHARS:
        return JsonResponse({
            "success": False,
            "error": f"Code checking is limited to {MAX_EXECUTION_CHARS} characters.",
        }, status=400)

    policy, _ = SandboxPolicy.objects.get_or_create(user=request.user)
    return JsonResponse(run_sandboxed_code(language, code, limits={
        "timeout_seconds": policy.timeout_seconds,
        "memory_mb": policy.memory_mb,
        "output_chars": policy.output_chars,
        "network_blocked": policy.network_blocked,
    }))


def _sandbox_policy_payload(policy):
    return {
        "timeout_seconds": policy.timeout_seconds,
        "memory_mb": policy.memory_mb,
        "output_chars": policy.output_chars,
        "require_approval": policy.require_approval,
        "network_blocked": policy.network_blocked,
    }


@login_required(login_url="/login/")
@require_http_methods(["GET", "POST"])
def sandbox_policy_api(request):
    policy, _ = SandboxPolicy.objects.get_or_create(user=request.user)
    if request.method == "POST":
        try:
            policy.timeout_seconds = min(10, max(1, int(request.POST.get("timeout_seconds", policy.timeout_seconds))))
            policy.memory_mb = min(256, max(32, int(request.POST.get("memory_mb", policy.memory_mb))))
            policy.output_chars = min(20000, max(1000, int(request.POST.get("output_chars", policy.output_chars))))
        except (TypeError, ValueError):
            return JsonResponse({"success": False, "error": "Runtime limits must be whole numbers."}, status=400)
        policy.require_approval = request.POST.get("require_approval", "true").lower() in {"1", "true", "yes", "on"}
        policy.network_blocked = True
        policy.save()
    return JsonResponse({"success": True, "policy": _sandbox_policy_payload(policy)})


@login_required(login_url="/login/")
@require_POST
def agent_plan(request):
    goal = request.POST.get("goal", "").strip()
    if not goal:
        return JsonResponse({"success": False, "error": "Describe the task for the agent."}, status=400)
    if len(goal) > 4000:
        return JsonResponse({"success": False, "error": "Agent goals are limited to 4,000 characters."}, status=400)
    task = AgentTask.objects.create(
        owner=request.user,
        title=goal[:80],
        goal=goal,
        plan=_agent_plan(goal),
    )
    return JsonResponse({"success": True, "task": _agent_payload(task)})


def _agent_team_payload(team):
    members = list(team.members or [])
    completed = sum(1 for member in members if member.get("status") == "completed")
    return {
        "id": team.id,
        "title": team.title,
        "goal": team.goal,
        "roles": team.roles,
        "members": members,
        "shared_context": team.shared_context,
        "status": team.status,
        "logs": team.logs,
        "current_member": team.current_member,
        "progress": completed,
        "total_members": len(members),
        "created_at": team.created_at.strftime("%d-%m-%Y %H:%M"),
    }


@login_required(login_url="/login/")
@require_http_methods(["GET", "POST"])
def agent_team_plan(request):
    if request.method == "GET":
        teams = AgentTeam.objects.filter(owner=request.user).order_by("-updated_at")[:30]
        return JsonResponse({"success": True, "teams": [_agent_team_payload(team) for team in teams]})
    goal = request.POST.get("goal", "").strip()
    if not goal:
        return JsonResponse({"success": False, "error": "Describe the team goal."}, status=400)
    try:
        roles = json.loads(request.POST.get("roles", "[]"))
    except json.JSONDecodeError:
        roles = []
    roles = [str(role).strip()[:40] for role in roles if str(role).strip()][:8]
    if not roles:
        roles = ["planner", "coder", "tester", "security", "documenter"]
    members = []
    task_ids = []
    for role in roles:
        child = AgentTask.objects.create(
            owner=request.user,
            title=f"{role.title()} · {goal[:120]}",
            goal=f"{role.title()} role for: {goal}",
            plan=_agent_plan(goal),
        )
        task_ids.append(child.id)
        members.append({"role": role, "task_id": child.id, "status": "queued", "result": ""})
    team = AgentTeam.objects.create(
        owner=request.user,
        title=goal[:80],
        goal=goal,
        roles=roles,
        members=members,
        shared_context={"child_task_ids": task_ids, "approval_required": True},
        logs=[{"level": "info", "message": "Agent team planned with approval-gated role tasks."}],
    )
    return JsonResponse({"success": True, "team": _agent_team_payload(team)}, status=201)


@login_required(login_url="/login/")
@require_POST
def agent_team_run(request, team_id):
    team = get_object_or_404(AgentTeam, id=team_id, owner=request.user)
    if team.status in {"completed", "blocked"}:
        return JsonResponse({"success": False, "error": "This agent team cannot run in its current state."}, status=400)
    members = list(team.members or [])
    if not members:
        return JsonResponse({"success": False, "error": "The agent team has no role tasks."}, status=400)
    team.status = "running"
    if team.current_member >= len(members):
        team.current_member = 0
    member = members[team.current_member]
    member["status"] = "running"
    member["result"] = "Role is active. Approve protected child steps before execution."
    team.logs = [*(team.logs or []), {"level": "info", "message": f"{member['role']} role started."}][-50:]
    team.save(update_fields=["status", "members", "logs", "updated_at"])
    return JsonResponse({"success": True, "team": _agent_team_payload(team)})


@login_required(login_url="/login/")
@require_POST
def agent_team_control(request, team_id):
    team = get_object_or_404(AgentTeam, id=team_id, owner=request.user)
    action = request.POST.get("action", "").strip().lower()
    if action == "pause":
        team.status = "paused"
        message = "Agent team paused by user."
    elif action == "resume":
        team.status = "running"
        message = "Agent team resumed. Review the active role approvals."
    elif action == "advance":
        members = list(team.members or [])
        if members and team.current_member < len(members):
            members[team.current_member]["status"] = "completed"
            members[team.current_member]["result"] = "Role checkpoint completed."
            team.current_member += 1
            team.members = members
        team.status = "completed" if team.current_member >= len(members) else "running"
        message = "Role checkpoint advanced."
    elif action == "cancel":
        team.status = "blocked"
        message = "Agent team cancelled by user."
    else:
        return JsonResponse({"success": False, "error": "Use pause, resume, advance, or cancel."}, status=400)
    team.logs = [*(team.logs or []), {"level": "warning" if action in {"pause", "cancel"} else "info", "message": message}][-50:]
    team.save(update_fields=["status", "members", "current_member", "logs", "updated_at"])
    return JsonResponse({"success": True, "team": _agent_team_payload(team)})


def _agent_job_payload(job):
    return {
        "id": job.id,
        "task_id": job.task_id,
        "task_title": job.task.title,
        "status": job.status,
        "payload": job.payload,
        "checkpoint": job.checkpoint,
        "logs": job.logs,
        "attempts": job.attempts,
        "last_error": job.last_error,
        "started_at": job.started_at.isoformat() if job.started_at else None,
        "finished_at": job.finished_at.isoformat() if job.finished_at else None,
        "created_at": job.created_at.isoformat(),
    }


@login_required(login_url="/login/")
@require_http_methods(["GET", "POST"])
def agent_jobs(request):
    if request.method == "GET":
        jobs = AgentJob.objects.filter(owner=request.user).select_related("task")[:100]
        return JsonResponse({"success": True, "jobs": [_agent_job_payload(job) for job in jobs]})
    task = get_object_or_404(AgentTask, id=request.POST.get("task_id"), owner=request.user)
    job = AgentJob.objects.create(
        owner=request.user,
        task=task,
        payload={"code": request.POST.get("code", "")[:20000], "test_code": request.POST.get("test_code", "")[:20000]},
        logs=[{"level": "info", "message": "Job queued and ready for a guarded worker run."}],
    )
    return JsonResponse({"success": True, "job": _agent_job_payload(job)}, status=201)


@login_required(login_url="/login/")
@require_POST
def agent_job_run(request, job_id):
    job = get_object_or_404(AgentJob, id=job_id, owner=request.user)
    if job.status in {"cancelled", "completed"}:
        return JsonResponse({"success": False, "error": "This job is no longer runnable."}, status=400)
    job.status = "running"
    job.attempts += 1
    job.started_at = job.started_at or timezone.now()
    job.checkpoint = {**(job.checkpoint or {}), "stage": "safe_checks", "updated_at": timezone.now().isoformat()}
    logs = [*(job.logs or []), {"level": "info", "message": f"Worker attempt {job.attempts} started."}]
    code = job.payload.get("code", "")
    test_code = job.payload.get("test_code", "")
    if code:
        scan = scan_files([{"filename": "background-agent-buffer", "content": code}])
        logs.append({"level": "warning" if scan["findings"] else "info", "message": scan["summary"]})
    if test_code:
        result = run_sandboxed_code("python", code + "\n\n" + test_code)
        logs.append({"level": "info" if result["success"] else "error", "message": "Guarded tests " + ("passed." if result["success"] else "failed.")})
        job.checkpoint["tests_passed"] = result["success"]
        if not result["success"]:
            job.status = "failed"
            job.last_error = result.get("stderr") or "Guarded tests failed."
    if job.status == "running":
        job.status = "completed"
        job.checkpoint["stage"] = "completed"
        job.finished_at = timezone.now()
        logs.append({"level": "info", "message": "Background job completed at a safe checkpoint."})
    job.logs = logs[-100:]
    job.save(update_fields=["status", "attempts", "started_at", "finished_at", "checkpoint", "logs", "last_error", "updated_at"])
    return JsonResponse({"success": job.status == "completed", "job": _agent_job_payload(job)})


@login_required(login_url="/login/")
@require_POST
def agent_job_control(request, job_id):
    job = get_object_or_404(AgentJob, id=job_id, owner=request.user)
    action = request.POST.get("action", "").strip().lower()
    if action == "pause":
        job.status = "paused"
        message = "Job paused at its last checkpoint."
    elif action == "resume":
        job.status = "queued"
        message = "Job resumed and queued for another worker run."
    elif action == "retry":
        job.status = "queued"
        job.last_error = ""
        message = "Job reset for retry."
    elif action == "cancel":
        job.status = "cancelled"
        message = "Job cancelled by user."
    else:
        return JsonResponse({"success": False, "error": "Use pause, resume, retry, or cancel."}, status=400)
    job.logs = [*(job.logs or []), {"level": "warning" if action in {"pause", "cancel"} else "info", "message": message}][-100:]
    job.save(update_fields=["status", "last_error", "logs", "updated_at"])
    return JsonResponse({"success": True, "job": _agent_job_payload(job)})


@login_required(login_url="/login/")
@require_http_methods(["GET", "POST"])
def agent_task(request, task_id):
    task = get_object_or_404(AgentTask, id=task_id, owner=request.user)
    if request.method == "GET":
        return JsonResponse({"success": True, "task": _agent_payload(task)})

    action = request.POST.get("action", "").strip().lower()
    try:
        step_index = int(request.POST.get("step", task.current_step))
    except (TypeError, ValueError):
        return JsonResponse({"success": False, "error": "Invalid plan step."}, status=400)
    if action == "undo":
        task.plan = _agent_plan(task.goal)
        task.current_step = 0
        task.status = "planned"
        task.result = "Task reset. No approved actions remain."
        task.control_state = "ready"
        task.logs = []
        task.retry_count = 0
    elif action in {"approve", "reject", "complete"}:
        if step_index < 0 or step_index >= len(task.plan):
            return JsonResponse({"success": False, "error": "Plan step does not exist."}, status=400)
        step = task.plan[step_index]
        if action == "approve":
            step["approved"] = True
            task.status = "running"
            task.current_step = step_index
            task.result = "Step approved. Review the proposed action before execution."
        elif action == "reject":
            step["approved"] = False
            task.status = "blocked"
            task.result = "Step rejected. The agent is paused until the plan is reset."
        else:
            if step.get("requires_approval") and not step.get("approved"):
                return JsonResponse({"success": False, "error": "Approve this step before marking it complete."}, status=400)
            step["completed"] = True
            task.current_step = min(step_index + 1, len(task.plan) - 1)
            task.status = "completed" if all(item.get("completed") for item in task.plan) else "running"
            task.result = "Step completed."
    else:
        return JsonResponse({"success": False, "error": "Use approve, reject, complete, or undo."}, status=400)
    task.save(update_fields=["plan", "current_step", "status", "result", "control_state", "logs", "retry_count", "updated_at"])
    return JsonResponse({"success": True, "task": _agent_payload(task)})


@login_required(login_url="/login/")
@require_POST
def agent_run(request, task_id):
    task = get_object_or_404(AgentTask, id=task_id, owner=request.user)
    if task.control_state == "cancelled":
        return JsonResponse({"success": False, "error": "This task was cancelled."}, status=400)
    if any(step.get("requires_approval") and not step.get("approved") for step in task.plan):
        return JsonResponse({"success": False, "error": "Approve protected plan steps before running the task."}, status=400)
    code = request.POST.get("code", "").strip()
    test_code = request.POST.get("test_code", "").strip()
    branch = request.POST.get("branch", "").strip()
    logs = list(task.logs or [])
    logs.append({"message": "Run started. Protected actions remain approval-gated.", "level": "info"})
    task.status = "running"
    task.control_state = "running"
    task.retry_count += 1
    if branch:
        task.source_branch = branch[:200]
        logs.append({"message": "Branch requested: " + task.source_branch, "level": "info"})
    if code:
        scan = scan_files([{"filename": "agent-buffer", "content": code}])
        logs.append({"message": scan["summary"], "level": "warning" if scan["findings"] else "info"})
    if test_code:
        result = run_sandboxed_code("python", code + "\n\n" + test_code)
        logs.append({"message": "Guarded tests " + ("passed." if result["success"] else "failed."), "level": "info" if result["success"] else "error"})
    logs.append({"message": "Run paused after safe checks. Apply file patches and Git writes explicitly.", "level": "info"})
    task.logs = logs[-50:]
    task.result = logs[-1]["message"]
    task.save(update_fields=["status", "control_state", "retry_count", "source_branch", "logs", "result", "updated_at"])
    return JsonResponse({"success": True, "task": _agent_payload(task)})


@login_required(login_url="/login/")
@require_POST
def agent_control(request, task_id):
    task = get_object_or_404(AgentTask, id=task_id, owner=request.user)
    action = request.POST.get("action", "").strip().lower()
    if action == "pause":
        task.control_state = "paused"
        task.result = "Task paused by user."
    elif action == "resume":
        task.control_state = "running"
        task.status = "running"
        task.result = "Task resumed. Review the next protected action."
    elif action == "retry":
        task.control_state = "ready"
        task.status = "planned"
        task.result = "Task ready for another approved run."
    elif action == "cancel":
        task.control_state = "cancelled"
        task.status = "blocked"
        task.result = "Task cancelled by user."
    else:
        return JsonResponse({"success": False, "error": "Use pause, resume, retry, or cancel."}, status=400)
    task.logs = [*(task.logs or []), {"message": task.result, "level": "warning" if action in {"pause", "cancel"} else "info"}][-50:]
    task.save(update_fields=["control_state", "status", "result", "logs", "updated_at"])
    return JsonResponse({"success": True, "task": _agent_payload(task)})


@login_required(login_url="/login/")
@require_GET
def ai_observability(request):
    events = list(AiEvent.objects.filter(owner=request.user).order_by("-created_at")[:100])
    tasks = list(AgentTask.objects.filter(owner=request.user).order_by("-updated_at")[:30])
    teams = list(AgentTeam.objects.filter(owner=request.user).order_by("-updated_at")[:20])
    jobs = list(AgentJob.objects.filter(owner=request.user).select_related("task").order_by("-updated_at")[:30])
    total = len(events)
    successful = sum(1 for event in events if event.success)
    durations = [event.duration_ms for event in events if event.duration_ms]
    model_metrics = {}
    for event in events:
        key = event.model_name or "unknown"
        row = model_metrics.setdefault(key, {"model": key, "requests": 0, "successful": 0, "duration_ms": 0, "output_chars": 0, "first_token_ms": []})
        row["requests"] += 1
        row["successful"] += 1 if event.success else 0
        row["duration_ms"] += event.duration_ms or 0
        row["output_chars"] += event.output_chars or 0
        first_token_ms = (event.metadata or {}).get("first_token_ms")
        if first_token_ms:
            row["first_token_ms"].append(int(first_token_ms))
    for row in model_metrics.values():
        row["success_rate"] = round((row["successful"] / row["requests"]) * 100, 1) if row["requests"] else 0
        row["average_duration_ms"] = round(row["duration_ms"] / row["requests"]) if row["requests"] else 0
        row["output_chars_per_second"] = round(row["output_chars"] / max(row["duration_ms"] / 1000, 0.001), 1)
        first_tokens = row.pop("first_token_ms")
        row["average_first_token_ms"] = round(sum(first_tokens) / len(first_tokens)) if first_tokens else 0
    timeline = []
    for task in tasks:
        for item in (task.logs or [])[-10:]:
            timeline.append({
                "kind": "task",
                "name": task.title,
                "status": task.status,
                "message": item.get("message", str(item)) if isinstance(item, dict) else str(item),
                "level": item.get("level", "info") if isinstance(item, dict) else "info",
                "created_at": task.updated_at.isoformat(),
            })
    for team in teams:
        for item in (team.logs or [])[-10:]:
            timeline.append({
                "kind": "team",
                "name": team.title,
                "status": team.status,
                "message": item.get("message", str(item)) if isinstance(item, dict) else str(item),
                "level": item.get("level", "info") if isinstance(item, dict) else "info",
                "created_at": team.updated_at.isoformat(),
            })
    for job in jobs:
        for item in (job.logs or [])[-10:]:
            timeline.append({
                "kind": "job",
                "name": job.task.title,
                "status": job.status,
                "message": item.get("message", str(item)) if isinstance(item, dict) else str(item),
                "level": item.get("level", "info") if isinstance(item, dict) else "info",
                "created_at": job.updated_at.isoformat(),
            })
    timeline.sort(key=lambda item: item["created_at"], reverse=True)
    return JsonResponse({
        "success": True,
        "summary": {
            "events": total,
            "successful": successful,
            "failed": total - successful,
            "success_rate": round((successful / total) * 100, 1) if total else 0,
            "average_duration_ms": round(sum(durations) / len(durations)) if durations else 0,
            "input_chars": sum(event.input_chars for event in events),
            "output_chars": sum(event.output_chars for event in events),
            "output_chars_per_second": round(sum(event.output_chars for event in events) / max(sum(durations) / 1000, 0.001), 1) if durations else 0,
            "average_first_token_ms": round(sum((event.metadata or {}).get("first_token_ms", 0) for event in events if (event.metadata or {}).get("first_token_ms") ) / max(sum(1 for event in events if (event.metadata or {}).get("first_token_ms")), 1)),
        },
        "agent_summary": {
            "tasks": len(tasks),
            "running_tasks": sum(1 for task in tasks if task.status == "running"),
            "teams": len(teams),
            "jobs": len(jobs),
            "active_jobs": sum(1 for job in jobs if job.status in {"queued", "running", "paused"}),
            "failed_jobs": sum(1 for job in jobs if job.status == "failed"),
            "timeline_events": len(timeline),
        },
        "resource_usage": {
            "cpu_ms": sum(int((event.metadata or {}).get("cpu_ms", 0) or 0) for event in events),
            "memory_mb_peak": max([int((event.metadata or {}).get("memory_mb", 0) or 0) for event in events] or [0]),
            "tool_calls": sum(int((event.metadata or {}).get("tool_calls", 0) or 0) for event in events),
        },
        "model_metrics": list(model_metrics.values()),
        "timeline": timeline[:120],
        "events": [
            {
                "id": event.id,
                "type": event.event_type,
                "model": event.model_name,
                "duration_ms": event.duration_ms,
                "input_chars": event.input_chars,
                "output_chars": event.output_chars,
                "success": event.success,
                "metadata": event.metadata,
                "created_at": event.created_at.isoformat(),
            }
            for event in events
        ],
    })


@login_required(login_url="/login/")
@require_GET
def agent_event_stream(request):
    def stream():
        latest_id = 0
        for _ in range(20):
            events = list(
                AiEvent.objects.filter(owner=request.user, id__gt=latest_id)
                .order_by("id")[:50]
            )
            for event in events:
                latest_id = max(latest_id, event.id)
                payload = {
                    "id": event.id,
                    "type": event.event_type,
                    "model": event.model_name,
                    "success": event.success,
                    "duration_ms": event.duration_ms,
                    "output_chars": event.output_chars,
                    "created_at": event.created_at.isoformat(),
                }
                yield f"event: ai\ndata: {json.dumps(payload)}\n\n"
            yield f": heartbeat {timezone.now().isoformat()}\n\n"
            time.sleep(1)

    response = StreamingHttpResponse(stream(), content_type="text/event-stream")
    response["Cache-Control"] = "no-cache"
    response["X-Accel-Buffering"] = "no"
    return response


def _evaluation_task_payload(task):
    return {
        "id": task.id,
        "name": task.name,
        "prompt": task.prompt,
        "code": task.code,
        "expected_output": task.expected_output,
        "language": task.language,
        "tags": [tag.strip() for tag in task.tags.split(",") if tag.strip()],
        "is_active": task.is_active,
        "created_at": task.created_at.isoformat(),
        "updated_at": task.updated_at.isoformat(),
    }


@login_required(login_url="/login/")
@require_http_methods(["GET", "POST"])
def evaluation_tasks(request):
    if request.method == "POST":
        name = request.POST.get("name", "").strip()[:160]
        prompt = request.POST.get("prompt", "").strip()
        if not name or not prompt:
            return JsonResponse({"success": False, "error": "Task name and prompt are required."}, status=400)
        task = EvaluationTask.objects.create(
            owner=request.user,
            name=name,
            prompt=prompt,
            code=request.POST.get("code", ""),
            expected_output=request.POST.get("expected_output", ""),
            language=request.POST.get("language", "auto").strip()[:50] or "auto",
            tags=request.POST.get("tags", "").strip()[:300],
        )
        return JsonResponse({"success": True, "task": _evaluation_task_payload(task)}, status=201)
    query = request.GET.get("q", "").strip()
    tasks = EvaluationTask.objects.filter(owner=request.user)
    if query:
        tasks = tasks.filter(Q(name__icontains=query) | Q(prompt__icontains=query) | Q(tags__icontains=query))
    return JsonResponse({"success": True, "tasks": [_evaluation_task_payload(task) for task in tasks[:100]]})


@login_required(login_url="/login/")
@require_http_methods(["POST", "DELETE"])
def evaluation_task_detail(request, task_id):
    task = get_object_or_404(EvaluationTask, id=task_id, owner=request.user)
    if request.method == "DELETE":
        task.delete()
        return JsonResponse({"success": True, "deleted": task_id})
    task.name = request.POST.get("name", task.name).strip()[:160] or task.name
    task.prompt = request.POST.get("prompt", task.prompt).strip() or task.prompt
    task.code = request.POST.get("code", task.code)
    task.expected_output = request.POST.get("expected_output", task.expected_output)
    task.language = request.POST.get("language", task.language).strip()[:50] or "auto"
    task.tags = request.POST.get("tags", task.tags).strip()[:300]
    task.is_active = request.POST.get("is_active", "true").lower() not in {"false", "0", "off"}
    task.save()
    return JsonResponse({"success": True, "task": _evaluation_task_payload(task)})


def _evaluation_run_payload(run):
    payload = {
        "id": run.id,
        "task_id": run.task_id,
        "model_name": run.model_name,
        "response": run.response,
        "status": run.status,
        "duration_ms": run.duration_ms,
        "input_chars": run.input_chars,
        "output_chars": run.output_chars,
        "metadata": run.metadata,
        "created_at": run.created_at.isoformat(),
    }
    score = getattr(run, "score", None)
    payload["score"] = _evaluation_score_payload(score) if score else None
    return payload


def _evaluation_score_payload(score):
    return {
        "correctness": score.correctness,
        "relevance": score.relevance,
        "completeness": score.completeness,
        "safety": score.safety,
        "overall": score.overall,
        "notes": score.notes,
        "method": score.method,
        "created_at": score.created_at.isoformat(),
    }


@login_required(login_url="/login/")
@require_POST
def evaluation_run(request, task_id):
    task = get_object_or_404(EvaluationTask, id=task_id, owner=request.user, is_active=True)
    settings = _user_ollama_settings(request.user)
    model_name = request.POST.get("model", "").strip()[:100] or settings.default_model
    evaluation_prompt = (
        "You are being evaluated as a local coding assistant.\n"
        f"Language: {task.language}\n"
        f"Task: {task.prompt}\n"
        f"Code fixture:\n{task.code}\n"
        "Return a practical answer with corrected code when appropriate."
    )
    run = EvaluationRun.objects.create(
        owner=request.user,
        task=task,
        model_name=model_name,
        input_chars=len(evaluation_prompt),
    )
    started = time.perf_counter()
    try:
        response = requests.post(
            f"{settings.server_url}/api/generate",
            json={
                "model": model_name,
                "prompt": evaluation_prompt,
                "stream": False,
                "options": {
                    "temperature": settings.temperature,
                    "top_p": settings.top_p,
                    "num_ctx": max(512, settings.max_context_chars // 4),
                },
            },
            timeout=(10, 300),
        )
        duration_ms = round((time.perf_counter() - started) * 1000)
        run.duration_ms = duration_ms
        if not response.ok:
            run.status = "failed"
            run.response = response.text[:4000]
            run.metadata = {"status_code": response.status_code}
        else:
            payload = response.json()
            run.response = str(payload.get("response", "")).strip()
            run.output_chars = len(run.response)
            run.status = "completed"
            run.metadata = {"done": payload.get("done", True)}
    except (requests.RequestException, ValueError) as exc:
        run.duration_ms = round((time.perf_counter() - started) * 1000)
        run.status = "failed"
        run.response = str(exc)
        run.metadata = {"error": str(exc)}
    run.save(update_fields=["status", "response", "duration_ms", "output_chars", "metadata"])
    AiEvent.objects.create(
        owner=request.user,
        event_type="evaluation",
        model_name=model_name,
        duration_ms=run.duration_ms,
        input_chars=run.input_chars,
        output_chars=run.output_chars,
        success=run.status == "completed",
        metadata={"task_id": task.id, "run_id": run.id},
    )
    return JsonResponse({"success": run.status == "completed", "run": _evaluation_run_payload(run)}, status=200 if run.status == "completed" else 502)


def _bounded_score(value):
    return min(max(int(value), 0), 100)


def _automatic_evaluation_scores(run):
    expected_tokens = set(re.findall(r"[a-z0-9_]+", run.task.expected_output.lower()))
    response_tokens = set(re.findall(r"[a-z0-9_]+", run.response.lower()))
    prompt_tokens = set(re.findall(r"[a-z0-9_]+", run.task.prompt.lower()))
    correctness = round(len(expected_tokens & response_tokens) / len(expected_tokens) * 100) if expected_tokens else (70 if run.response else 0)
    relevance = round(len(prompt_tokens & response_tokens) / len(prompt_tokens) * 100) if prompt_tokens else 0
    completeness = 85 if len(run.response) >= 80 else (60 if run.response else 0)
    safety = 100 if not re.search(r"(api[_ -]?key|password|secret|rm\s+-rf)", run.response, re.IGNORECASE) else 35
    overall = round((correctness + relevance + completeness + safety) / 4)
    return correctness, relevance, completeness, safety, overall


@login_required(login_url="/login/")
@require_POST
def evaluation_score(request, run_id):
    run = get_object_or_404(EvaluationRun, id=run_id, owner=request.user)
    automatic = request.POST.get("automatic", "true").lower() not in {"false", "0", "off"}
    if automatic:
        values = _automatic_evaluation_scores(run)
        method = "automatic"
    else:
        try:
            values = tuple(_bounded_score(request.POST.get(field, 0)) for field in ("correctness", "relevance", "completeness", "safety"))
        except (TypeError, ValueError):
            return JsonResponse({"success": False, "error": "Scores must be whole numbers from 0 to 100."}, status=400)
        values = (*values, round(sum(values) / 4))
        method = "reviewer"
    score, _ = EvaluationScore.objects.update_or_create(
        run=run,
        defaults={
            "correctness": values[0],
            "relevance": values[1],
            "completeness": values[2],
            "safety": values[3],
            "overall": values[4],
            "notes": request.POST.get("notes", "").strip()[:4000],
            "method": method,
        },
    )
    return JsonResponse({"success": True, "score": _evaluation_score_payload(score), "run": _evaluation_run_payload(run)})


def _regression_suite_payload(suite):
    return {
        "id": suite.id,
        "name": suite.name,
        "description": suite.description,
        "task_ids": suite.task_ids,
        "baseline": suite.baseline,
        "last_result": suite.last_result,
        "last_run_at": suite.last_run_at.isoformat() if suite.last_run_at else None,
    }


@login_required(login_url="/login/")
@require_http_methods(["GET", "POST"])
def evaluation_regressions(request):
    if request.method == "POST":
        name = request.POST.get("name", "").strip()[:160]
        if not name:
            return JsonResponse({"success": False, "error": "Regression suite name is required."}, status=400)
        try:
            requested_ids = json.loads(request.POST.get("task_ids", "[]"))
            task_ids = [int(value) for value in requested_ids]
        except (TypeError, ValueError, json.JSONDecodeError):
            return JsonResponse({"success": False, "error": "Task IDs must be a JSON array."}, status=400)
        valid_ids = set(EvaluationTask.objects.filter(owner=request.user, is_active=True, id__in=task_ids).values_list("id", flat=True))
        if not task_ids:
            task_ids = list(EvaluationTask.objects.filter(owner=request.user, is_active=True).values_list("id", flat=True))
        elif set(task_ids) != valid_ids:
            return JsonResponse({"success": False, "error": "Every task must belong to your account."}, status=400)
        suite = EvaluationRegressionSuite.objects.create(
            owner=request.user,
            name=name,
            description=request.POST.get("description", "").strip()[:1000],
            task_ids=task_ids,
        )
        return JsonResponse({"success": True, "suite": _regression_suite_payload(suite)}, status=201)
    suites = EvaluationRegressionSuite.objects.filter(owner=request.user)
    return JsonResponse({"success": True, "suites": [_regression_suite_payload(suite) for suite in suites[:100]]})


@login_required(login_url="/login/")
@require_POST
def evaluation_regression_run(request, suite_id):
    suite = get_object_or_404(EvaluationRegressionSuite, id=suite_id, owner=request.user)
    model_name = request.POST.get("model", "").strip()
    results = []
    baseline = dict(suite.baseline or {})
    if model_name:
        runs = EvaluationRun.objects.filter(owner=request.user, model_name=model_name, status="completed", score__isnull=False)
    else:
        runs = EvaluationRun.objects.filter(owner=request.user, status="completed", score__isnull=False)
    for task_id in suite.task_ids:
        run = runs.filter(task_id=task_id).select_related("score").order_by("-created_at").first()
        if not run:
            results.append({"task_id": task_id, "status": "missing", "score": None})
            continue
        current = run.score.overall
        key = str(task_id)
        if key not in baseline:
            baseline[key] = current
        threshold = max(0, int(baseline[key]) - 10)
        results.append({
            "task_id": task_id,
            "run_id": run.id,
            "model_name": run.model_name,
            "score": current,
            "baseline": baseline[key],
            "delta": current - int(baseline[key]),
            "status": "passed" if current >= threshold else "regressed",
        })
    passed = bool(results) and all(item["status"] == "passed" for item in results)
    suite.baseline = baseline
    suite.last_result = {"passed": passed, "model_name": model_name or "latest", "results": results}
    suite.last_run_at = timezone.now()
    suite.save(update_fields=["baseline", "last_result", "last_run_at", "updated_at"])
    return JsonResponse({"success": True, "suite": _regression_suite_payload(suite)})


@login_required(login_url="/login/")
@require_GET
def evaluation_dashboard(request):
    tasks = EvaluationTask.objects.filter(owner=request.user)
    runs = list(EvaluationRun.objects.filter(owner=request.user).select_related("task", "score").order_by("-created_at")[:200])
    completed = [run for run in runs if run.status == "completed"]
    scored = [run for run in completed if hasattr(run, "score")]
    durations = [run.duration_ms for run in completed if run.duration_ms]
    models = {}
    for run in runs:
        item = models.setdefault(run.model_name, {"model_name": run.model_name, "runs": 0, "completed": 0, "scores": [], "durations": []})
        item["runs"] += 1
        if run.status == "completed":
            item["completed"] += 1
            if run.duration_ms:
                item["durations"].append(run.duration_ms)
        if hasattr(run, "score"):
            item["scores"].append(run.score.overall)
    model_rows = []
    for item in models.values():
        model_rows.append({
            "model_name": item["model_name"],
            "runs": item["runs"],
            "completed": item["completed"],
            "average_score": round(sum(item["scores"]) / len(item["scores"]), 1) if item["scores"] else 0,
            "average_duration_ms": round(sum(item["durations"]) / len(item["durations"])) if item["durations"] else 0,
        })
    suites = list(EvaluationRegressionSuite.objects.filter(owner=request.user))
    return JsonResponse({
        "success": True,
        "summary": {
            "tasks": tasks.count(),
            "runs": len(runs),
            "completed_runs": len(completed),
            "success_rate": round((len(completed) / len(runs)) * 100, 1) if runs else 0,
            "scored_runs": len(scored),
            "average_score": round(sum(run.score.overall for run in scored) / len(scored), 1) if scored else 0,
            "average_duration_ms": round(sum(durations) / len(durations)) if durations else 0,
            "regression_suites": len(suites),
            "regressions_needing_review": sum(1 for suite in suites if suite.last_result and not suite.last_result.get("passed", False)),
        },
        "models": sorted(model_rows, key=lambda row: (-row["average_score"], row["model_name"])),
        "recent_runs": [
            {
                "id": run.id,
                "task": run.task.name,
                "model_name": run.model_name,
                "status": run.status,
                "duration_ms": run.duration_ms,
                "score": run.score.overall if hasattr(run, "score") else None,
                "created_at": run.created_at.isoformat(),
            }
            for run in runs[:30]
        ],
    })


def _workspace_for_user(user):
    membership = (
        WorkspaceMembership.objects.filter(user=user)
        .select_related("workspace")
        .order_by("workspace_id")
        .first()
    )
    if membership:
        return membership.workspace
    workspace = Workspace.objects.create(owner=user, name=f"{user.username}'s workspace")
    WorkspaceMembership.objects.create(workspace=workspace, user=user, role="admin")
    AuditEvent.objects.create(
        actor=user,
        workspace=workspace,
        event_type="workspace.created",
        details={"name": workspace.name},
    )
    return workspace


def _workspace_membership(user, workspace):
    return WorkspaceMembership.objects.filter(user=user, workspace=workspace).first()


def _workspace_payload(workspace):
    members = (
        WorkspaceMembership.objects.filter(workspace=workspace)
        .select_related("user")
        .order_by("user__username")
    )
    audits = (
        AuditEvent.objects.filter(workspace=workspace)
        .select_related("actor")
        .order_by("-created_at")[:50]
    )
    return {
        "id": workspace.id,
        "name": workspace.name,
        "owner": workspace.owner.username,
        "members": [
            {"id": member.id, "username": member.user.username, "role": member.role}
            for member in members
        ],
        "audit": [
            {
                "event_type": event.event_type,
                "actor": event.actor.username if event.actor else "system",
                "details": event.details,
                "created_at": event.created_at.isoformat(),
            }
            for event in audits
        ],
    }


@login_required(login_url="/login/")
@require_http_methods(["GET", "POST"])
def workspace_api(request):
    workspace = None
    workspace_id = (
        request.POST.get("workspace_id")
        or request.GET.get("workspace_id")
        or request.session.get("workspace_id")
    )
    if workspace_id:
        try:
            candidate = get_object_or_404(Workspace, id=int(workspace_id))
            if _workspace_membership(request.user, candidate):
                workspace = candidate
        except (TypeError, ValueError):
            workspace = None
    if workspace is None:
        workspace = _workspace_for_user(request.user)
    membership = _workspace_membership(request.user, workspace)
    if not membership:
        return JsonResponse({"success": False, "error": "You are not a workspace member."}, status=403)
    request.session["workspace_id"] = workspace.id

    if request.method == "POST":
        action = request.POST.get("action", "").strip().lower()
        if action == "create":
            name = request.POST.get("name", "").strip()[:120]
            if not name:
                return JsonResponse({"success": False, "error": "Workspace name is required."}, status=400)
            workspace = Workspace.objects.create(owner=request.user, name=name)
            WorkspaceMembership.objects.create(workspace=workspace, user=request.user, role="admin")
            AuditEvent.objects.create(actor=request.user, workspace=workspace, event_type="workspace.created", details={"name": name})
            request.session["workspace_id"] = workspace.id
        elif action == "add_member":
            if membership.role != "admin":
                return JsonResponse({"success": False, "error": "Only workspace admins can manage members."}, status=403)
            username = request.POST.get("username", "").strip()
            role = request.POST.get("role", "developer").strip()
            if role not in dict(WorkspaceMembership.ROLE_CHOICES):
                return JsonResponse({"success": False, "error": "Invalid workspace role."}, status=400)
            target = get_object_or_404(User, username=username)
            member, _ = WorkspaceMembership.objects.update_or_create(
                workspace=workspace,
                user=target,
                defaults={"role": role},
            )
            AuditEvent.objects.create(
                actor=request.user,
                workspace=workspace,
                event_type="member.role_changed",
                details={"username": target.username, "role": member.role},
            )
        elif action == "set_role":
            if membership.role != "admin":
                return JsonResponse({"success": False, "error": "Only workspace admins can manage roles."}, status=403)
            member = get_object_or_404(WorkspaceMembership, id=request.POST.get("member_id"), workspace=workspace)
            role = request.POST.get("role", "").strip()
            if role not in dict(WorkspaceMembership.ROLE_CHOICES):
                return JsonResponse({"success": False, "error": "Invalid workspace role."}, status=400)
            member.role = role
            member.save(update_fields=["role"])
            AuditEvent.objects.create(actor=request.user, workspace=workspace, event_type="member.role_changed", details={"username": member.user.username, "role": role})
        elif action == "remove_member":
            if membership.role != "admin":
                return JsonResponse({"success": False, "error": "Only workspace admins can remove members."}, status=403)
            member = get_object_or_404(WorkspaceMembership, id=request.POST.get("member_id"), workspace=workspace)
            if member.user_id == workspace.owner_id:
                return JsonResponse({"success": False, "error": "The workspace owner cannot be removed."}, status=400)
            username = member.user.username
            member.delete()
            AuditEvent.objects.create(actor=request.user, workspace=workspace, event_type="member.removed", details={"username": username})
        elif action not in {"", "switch"}:
            return JsonResponse({"success": False, "error": "Unsupported workspace action."}, status=400)

    return JsonResponse({"success": True, "workspace": _workspace_payload(workspace)})


def _policy_payload(policy, membership):
    return {
        "identity": {
            "username": membership.user.username,
            "role": membership.role,
            "authentication": "local Django account",
        },
        "policy": {
            "require_approval_for_git": policy.require_approval_for_git,
            "require_approval_for_tools": policy.require_approval_for_tools,
            "require_approval_for_deploy": policy.require_approval_for_deploy,
            "require_tests": policy.require_tests,
            "allow_external_connectors": policy.allow_external_connectors,
            "audit_retention_days": policy.audit_retention_days,
        },
    }


@login_required(login_url="/login/")
@require_http_methods(["GET", "POST"])
def identity_policy_api(request):
    workspace = _workspace_for_user(request.user)
    membership = _workspace_membership(request.user, workspace)
    if not membership:
        return JsonResponse({"success": False, "error": "You are not a workspace member."}, status=403)
    policy, _ = WorkspacePolicy.objects.get_or_create(workspace=workspace)
    if request.method == "POST":
        if membership.role != "admin":
            return JsonResponse({"success": False, "error": "Only workspace admins can change policy."}, status=403)
        for field in (
            "require_approval_for_git",
            "require_approval_for_tools",
            "require_approval_for_deploy",
            "require_tests",
            "allow_external_connectors",
        ):
            setattr(policy, field, request.POST.get(field, "false").lower() in {"1", "true", "yes", "on"})
        try:
            policy.audit_retention_days = min(3650, max(7, int(request.POST.get("audit_retention_days", "90"))))
        except (TypeError, ValueError):
            return JsonResponse({"success": False, "error": "Audit retention must be a number of days."}, status=400)
        policy.save()
        AuditEvent.objects.create(
            actor=request.user,
            workspace=workspace,
            event_type="policy.updated",
            details=_policy_payload(policy, membership)["policy"],
        )
    return JsonResponse({"success": True, "workspace": workspace.name, **_policy_payload(policy, membership)})


def _enterprise_identity_payload(config):
    return {
        "provider": config.provider,
        "issuer_url": config.issuer_url,
        "saml_entrypoint_url": config.saml_entrypoint_url,
        "client_id": config.client_id,
        "allowed_domains": config.allowed_domains,
        "enforce_sso": config.enforce_sso,
        "scim_enabled": config.scim_enabled,
        "token_configured": bool(config.scim_token_hash),
        "updated_at": config.updated_at.isoformat() if config.updated_at else None,
    }


@login_required(login_url="/login/")
@require_http_methods(["GET", "POST"])
def enterprise_identity_api(request):
    workspace = _workspace_for_user(request.user)
    membership = _workspace_membership(request.user, workspace)
    if not membership:
        return JsonResponse({"success": False, "error": "You are not a workspace member."}, status=403)
    config, _ = EnterpriseIdentityConfig.objects.get_or_create(workspace=workspace)
    scim_token = None
    if request.method == "POST":
        if membership.role != "admin":
            return JsonResponse({"success": False, "error": "Only workspace admins can change enterprise identity."}, status=403)
        provider = request.POST.get("provider", "oidc").strip().lower()
        if provider not in dict(EnterpriseIdentityConfig.PROVIDER_CHOICES):
            return JsonResponse({"success": False, "error": "Provider must be OIDC or SAML."}, status=400)
        config.provider = provider
        config.issuer_url = request.POST.get("issuer_url", "").strip()[:500]
        config.saml_entrypoint_url = request.POST.get("saml_entrypoint_url", "").strip()[:500]
        config.client_id = request.POST.get("client_id", "").strip()[:200]
        config.allowed_domains = request.POST.get("allowed_domains", "").strip()[:500]
        config.enforce_sso = request.POST.get("enforce_sso", "false").lower() in {"1", "true", "yes", "on"}
        config.scim_enabled = request.POST.get("scim_enabled", "false").lower() in {"1", "true", "yes", "on"}
        if request.POST.get("action", "").strip().lower() == "rotate_scim_token" or not config.scim_token_hash:
            scim_token = secrets.token_urlsafe(32)
            config.scim_token_hash = hashlib.sha256(scim_token.encode("utf-8")).hexdigest()
        config.save()
        AuditEvent.objects.create(
            actor=request.user,
            workspace=workspace,
            event_type="identity.enterprise_updated",
            details={"provider": config.provider, "scim_enabled": config.scim_enabled, "token_rotated": bool(scim_token)},
        )
    payload = {"success": True, "workspace": workspace.name, "identity_config": _enterprise_identity_payload(config)}
    if scim_token:
        payload["scim_token"] = scim_token
    return JsonResponse(payload)


def _secret_payload(item):
    return {
        "id": item.id,
        "name": item.name,
        "description": item.description,
        "version": item.version,
        "created_at": item.created_at.isoformat(),
        "updated_at": item.updated_at.isoformat(),
        "masked": True,
    }


@login_required(login_url="/login/")
@require_http_methods(["GET", "POST"])
def secrets_api(request):
    workspace = _workspace_for_user(request.user)
    membership = _workspace_membership(request.user, workspace)
    if not membership:
        return JsonResponse({"success": False, "error": "You are not a workspace member."}, status=403)
    if request.method == "POST":
        if membership.role != "admin":
            return JsonResponse({"success": False, "error": "Only workspace admins can manage secrets."}, status=403)
        name = request.POST.get("name", "").strip()[:120]
        value = request.POST.get("value", "")
        description = request.POST.get("description", "").strip()[:300]
        if not name or not re.fullmatch(r"[A-Za-z0-9_.-]+", name):
            return JsonResponse({"success": False, "error": "Secret names may contain letters, numbers, dots, underscores, and hyphens."}, status=400)
        if not value:
            return JsonResponse({"success": False, "error": "Secret value is required."}, status=400)
        item = SecretVaultItem.objects.filter(workspace=workspace, name=name).first()
        version = item.version + 1 if item else 1
        if item:
            item.ciphertext = encrypt_secret(value)
            item.description = description
            item.version = version
            item.save()
        else:
            item = SecretVaultItem.objects.create(workspace=workspace, name=name, description=description, ciphertext=encrypt_secret(value))
        AuditEvent.objects.create(actor=request.user, workspace=workspace, event_type="secret.updated", details={"name": name, "version": version})
        return JsonResponse({"success": True, "secret": _secret_payload(item)})
    return JsonResponse({"success": True, "secrets": [_secret_payload(item) for item in SecretVaultItem.objects.filter(workspace=workspace)]})


@login_required(login_url="/login/")
@require_POST
def secret_reveal(request, secret_id):
    workspace = _workspace_for_user(request.user)
    membership = _workspace_membership(request.user, workspace)
    if not membership or membership.role != "admin":
        return JsonResponse({"success": False, "error": "Only workspace admins can reveal secrets."}, status=403)
    item = get_object_or_404(SecretVaultItem, id=secret_id, workspace=workspace)
    try:
        value = decrypt_secret(item.ciphertext)
    except Exception:
        return JsonResponse({"success": False, "error": "Secret could not be decrypted."}, status=500)
    AuditEvent.objects.create(actor=request.user, workspace=workspace, event_type="secret.revealed", details={"name": item.name})
    return JsonResponse({"success": True, "name": item.name, "value": value, "warning": "Keep this value private; it will not be shown in vault listings."})


@login_required(login_url="/login/")
@require_http_methods(["DELETE"])
def secret_delete(request, secret_id):
    workspace = _workspace_for_user(request.user)
    membership = _workspace_membership(request.user, workspace)
    if not membership or membership.role != "admin":
        return JsonResponse({"success": False, "error": "Only workspace admins can delete secrets."}, status=403)
    item = get_object_or_404(SecretVaultItem, id=secret_id, workspace=workspace)
    name = item.name
    item.delete()
    AuditEvent.objects.create(actor=request.user, workspace=workspace, event_type="secret.deleted", details={"name": name})
    return JsonResponse({"success": True, "deleted_id": secret_id})


def _extension_payload(package, install=None):
    return {
        "slug": package.slug,
        "name": package.name,
        "version": package.version,
        "description": package.description,
        "permissions": package.permissions,
        "manifest": package.manifest,
        "installed": bool(install and install.status == "installed"),
        "install_status": install.status if install else "not_installed",
        "installed_version": install.version if install else None,
        "can_rollback": bool(install and install.previous_version),
    }


def _sync_extension_catalog():
    for item in EXTENSION_CATALOG:
        ExtensionPackage.objects.update_or_create(slug=item["slug"], defaults=item)


@login_required(login_url="/login/")
@require_http_methods(["GET", "POST"])
def extension_marketplace(request):
    _sync_extension_catalog()
    if request.method == "POST":
        action = request.POST.get("action", "install").strip().lower()
        if action == "publish":
            name = request.POST.get("name", "").strip()[:160]
            raw_slug = request.POST.get("slug", "").strip().lower()
            slug = re.sub(r"[^a-z0-9-]", "-", raw_slug).strip("-")[:80]
            slug = "custom-" + request.user.username.lower() + "-" + slug
            if not name or not slug or slug.endswith("-"):
                return JsonResponse({"success": False, "error": "Extension name and slug are required."}, status=400)
            try:
                permissions = json.loads(request.POST.get("permissions", "[]"))
                manifest = json.loads(request.POST.get("manifest", "{}"))
            except (TypeError, ValueError, json.JSONDecodeError):
                return JsonResponse({"success": False, "error": "Permissions and manifest must be valid JSON."}, status=400)
            if not isinstance(permissions, list) or not isinstance(manifest, dict):
                return JsonResponse({"success": False, "error": "Permissions must be a list and manifest must be an object."}, status=400)
            package, _ = ExtensionPackage.objects.update_or_create(
                slug=slug,
                defaults={
                    "owner": request.user,
                    "name": name,
                    "version": request.POST.get("version", "1.0.0").strip()[:40],
                    "description": request.POST.get("description", "").strip(),
                    "permissions": permissions[:20],
                    "manifest": manifest,
                },
            )
            return JsonResponse({"success": True, "extension": _extension_payload(package)})
        slug = request.POST.get("slug", "").strip()
        package = get_object_or_404(ExtensionPackage, slug=slug, is_active=True)
        install = ExtensionInstall.objects.filter(owner=request.user, package=package).first()
        if action == "install":
            try:
                approved = json.loads(request.POST.get("approved_permissions", "[]"))
            except (TypeError, ValueError, json.JSONDecodeError):
                approved = []
            approved = approved if isinstance(approved, list) else []
            required = set(package.permissions)
            if request.POST.get("permissions_approved") != "true" or not required.issubset(set(approved)):
                return JsonResponse({"success": False, "error": "Review and approve every requested permission before installation.", "required_permissions": package.permissions}, status=400)
            previous = install.version if install and install.version != package.version else ""
            install, _ = ExtensionInstall.objects.update_or_create(
                owner=request.user,
                package=package,
                defaults={"status": "installed", "version": package.version, "previous_version": previous, "approved_permissions": approved},
            )
        elif action == "uninstall":
            if not install:
                return JsonResponse({"success": False, "error": "Extension is not installed."}, status=404)
            install.status = "disabled"
            install.save(update_fields=["status", "updated_at"])
        elif action == "rollback":
            if not install or not install.previous_version:
                return JsonResponse({"success": False, "error": "No previous extension version is available."}, status=400)
            install.version, install.previous_version = install.previous_version, install.version
            install.status = "installed"
            install.save(update_fields=["version", "previous_version", "status", "updated_at"])
        else:
            return JsonResponse({"success": False, "error": "Use install, uninstall, rollback, or publish."}, status=400)
        return JsonResponse({"success": True, "extension": _extension_payload(package, install)})
    packages = []
    for package in ExtensionPackage.objects.filter(is_active=True).order_by("name"):
        packages.append(_extension_payload(package, ExtensionInstall.objects.filter(owner=request.user, package=package).first()))
    return JsonResponse({"success": True, "extensions": packages, "sdk": {"permissions": ["project.read", "agent.read", "agent.approve", "index.write"], "manifest_fields": ["name", "slug", "version", "description", "permissions", "manifest"]}})


def _scim_workspace(request):
    workspace_id = request.headers.get("X-Workspace-ID") or request.GET.get("workspace_id")
    if workspace_id:
        try:
            return Workspace.objects.get(id=int(workspace_id))
        except (TypeError, ValueError, Workspace.DoesNotExist):
            return None
    return Workspace.objects.order_by("id").first()


def _scim_payload(request):
    authorization = request.headers.get("Authorization", "")
    if not authorization.lower().startswith("bearer "):
        return None, "Bearer token required."
    token = authorization[7:].strip()
    workspace = _scim_workspace(request)
    if not workspace:
        return None, "Workspace not found."
    config = EnterpriseIdentityConfig.objects.filter(workspace=workspace).first()
    if not config or not config.scim_enabled or not config.scim_token_hash:
        return None, "SCIM is not enabled for this workspace."
    expected = hashlib.sha256(token.encode("utf-8")).hexdigest()
    if not secrets.compare_digest(expected, config.scim_token_hash):
        return None, "Invalid SCIM token."
    return workspace, None


@csrf_exempt
@require_http_methods(["GET", "POST", "DELETE"])
def scim_directory_api(request):
    workspace, error = _scim_payload(request)
    if error:
        return JsonResponse({"success": False, "error": error}, status=401)
    if request.method == "GET":
        members = WorkspaceMembership.objects.filter(workspace=workspace).select_related("user").order_by("user__username")
        resources = [
            {
                "id": str(member.user_id),
                "userName": member.user.username,
                "active": member.user.is_active,
                "displayName": member.user.get_full_name() or member.user.username,
                "role": member.role,
            }
            for member in members
        ]
        return JsonResponse({"schemas": ["urn:ietf:params:scim:api:messages:2.0:ListResponse"], "totalResults": len(resources), "Resources": resources})
    try:
        body = json.loads(request.body.decode("utf-8") or "{}")
    except (UnicodeDecodeError, json.JSONDecodeError):
        return JsonResponse({"success": False, "error": "SCIM request body must be valid JSON."}, status=400)
    username = (body.get("userName") or body.get("username") or request.GET.get("userName") or "").strip()[:150]
    if not username:
        return JsonResponse({"success": False, "error": "SCIM userName is required."}, status=400)
    if request.method == "DELETE":
        user = User.objects.filter(username=username).first()
        if not user:
            return JsonResponse({"success": False, "error": "User not found."}, status=404)
        membership = WorkspaceMembership.objects.filter(workspace=workspace, user=user).first()
        user.is_active = False
        user.save(update_fields=["is_active"])
        if membership and user.id != workspace.owner_id:
            membership.delete()
        DirectoryProvisioningEvent.objects.create(workspace=workspace, username=username, action="deprovisioned", details={"removed_membership": bool(membership and user.id != workspace.owner_id)})
        return JsonResponse({"success": True, "action": "deprovisioned", "userName": username})
    email = ""
    emails = body.get("emails") or []
    if isinstance(emails, list) and emails and isinstance(emails[0], dict):
        email = str(emails[0].get("value") or "").strip()[:254]
    user, created = User.objects.get_or_create(username=username, defaults={"email": email})
    if email and user.email != email:
        user.email = email
    active = body.get("active", True) is not False
    user.is_active = active
    display_name = str(body.get("displayName") or "").strip()
    if display_name:
        parts = display_name.split(" ", 1)
        user.first_name = parts[0][:150]
        user.last_name = parts[1][:150] if len(parts) > 1 else ""
    user.save()
    role = str(body.get("role") or "developer").strip().lower()
    if role not in dict(WorkspaceMembership.ROLE_CHOICES) or role == "admin":
        role = "developer"
    membership, _ = WorkspaceMembership.objects.update_or_create(workspace=workspace, user=user, defaults={"role": role})
    action = "provisioned" if created else "updated"
    DirectoryProvisioningEvent.objects.create(workspace=workspace, username=username, action=action, details={"active": active, "role": membership.role})
    return JsonResponse({"success": True, "action": action, "user": {"id": str(user.id), "userName": user.username, "active": user.is_active, "role": membership.role}}, status=201 if created else 200)


@login_required(login_url="/login/")
@require_POST
def devops_generate(request):
    kind = request.POST.get("kind", "dockerfile").strip()
    code = request.POST.get("code", "")[:30_000]
    filenames = request.POST.getlist("filenames")
    try:
        artifact = generate_artifact(kind, code, filenames)
    except ValueError as exc:
        return JsonResponse({"success": False, "error": str(exc)}, status=400)
    return JsonResponse({
        "success": True,
        "kind": kind,
        "artifact": artifact,
        "checklist": [
            "Review exposed ports and runtime command.",
            "Move secrets to environment variables or a secret manager.",
            "Run the generated CI job and security scan before deployment.",
        ],
    })


@login_required(login_url="/login/")
@require_POST
def devops_logs(request):
    logs = request.POST.get("logs", "")[:50_000]
    return JsonResponse({"success": True, **analyze_logs(logs)})


@login_required(login_url="/login/")
@require_POST
def deployment_kit(request):
    try:
        kit = generate_deployment_kit(
            project_name=request.POST.get("project_name", "syntax-local"),
            port=request.POST.get("port", "8000"),
            health_path=request.POST.get("health_path", "/"),
            strategy=request.POST.get("strategy", "rolling"),
        )
    except ValueError as exc:
        return JsonResponse({"success": False, "error": str(exc)}, status=400)
    return JsonResponse({"success": True, **kit})


@login_required(login_url="/login/")
@require_POST
def deployment_validate(request):
    return JsonResponse({"success": True, **validate_deployment_environment(request.POST.get("content", "")[:50_000])})


@login_required(login_url="/login/")
@require_POST
def documentation_generate(request):
    doc_type = request.POST.get("doc_type", "readme").strip()
    project_name = request.POST.get("project_name", "Local project")
    code = request.POST.get("code", "")[:30_000]
    filename = request.POST.get("filename", "current-code")
    change_summary = request.POST.get("change_summary", "")[:2_000]
    files = [{"filename": filename, "content": code}]
    raw_files = request.POST.get("files_json", "")
    if raw_files:
        try:
            parsed_files = json.loads(raw_files)
            if isinstance(parsed_files, list):
                files.extend(item for item in parsed_files[:100] if isinstance(item, dict))
        except (TypeError, ValueError):
            return JsonResponse({"success": False, "error": "files_json must be a JSON array."}, status=400)
    if doc_type not in {"readme", "api", "changelog", "onboarding"}:
        return JsonResponse({"success": False, "error": "Unsupported documentation type."}, status=400)
    markdown = generate_documentation(files, doc_type, project_name, change_summary)
    return JsonResponse({"success": True, "doc_type": doc_type, "markdown": markdown[:60_000]})


@login_required(login_url="/login/")
@require_POST
def review_gate_api(request):
    code = request.POST.get("code", "")[:30_000]
    filename = request.POST.get("filename", "editor-buffer")
    language = request.POST.get("language", "auto").strip().lower()
    diff = request.POST.get("diff", "")
    tests = request.POST.get("tests", "")
    files = [{"filename": filename, "content": code}] if code else []
    raw_files = request.POST.get("files_json", "")
    if raw_files:
        try:
            parsed = json.loads(raw_files)
            if not isinstance(parsed, list) or len(parsed) > MAX_PROJECT_FILES:
                raise ValueError("invalid file list")
            files = [
                {"filename": _safe_filename(item.get("filename", "uploaded-file")), "content": str(item.get("content", ""))}
                for item in parsed if isinstance(item, dict)
            ]
        except (TypeError, ValueError, json.JSONDecodeError):
            return JsonResponse({"success": False, "error": "Invalid review file list."}, status=400)
    if not files:
        return JsonResponse({"success": False, "error": "Provide code or project files to review."}, status=400)
    return JsonResponse({"success": True, **review_gate(files, language, diff, tests)})


@login_required(login_url="/login/")
@require_POST
def dependencies_analyze(request):
    filename = request.POST.get("filename", "requirements.txt")
    content = request.POST.get("content", "")[:30_000]
    files = [{"filename": filename, "content": content}]
    raw_files = request.POST.get("files_json", "")
    if raw_files:
        try:
            parsed = json.loads(raw_files)
            if not isinstance(parsed, list) or len(parsed) > MAX_PROJECT_FILES:
                raise ValueError("invalid dependency file list")
            files = [
                {"filename": _safe_filename(item.get("filename", "manifest")), "content": str(item.get("content", ""))}
                for item in parsed if isinstance(item, dict)
            ]
        except (TypeError, ValueError, json.JSONDecodeError):
            return JsonResponse({"success": False, "error": "Invalid dependency file list."}, status=400)
    return JsonResponse({"success": True, **analyze_dependencies(files)})


@login_required(login_url="/login/")
@require_POST
def browser_test_generate(request):
    return JsonResponse({
        "success": True,
        "test_code": generate_playwright_test(
            request.POST.get("base_url"),
            request.POST.get("flow"),
            request.POST.get("snapshot_name", "home"),
        ),
        "command": "npx playwright test --update-snapshots",
    })


@login_required(login_url="/login/")
@require_POST
def browser_test_report(request):
    return JsonResponse({"success": True, **analyze_browser_report(request.POST.get("report", "")[:50_000])})


@login_required(login_url="/login/")
@require_POST
def model_route(request):
    raw_models = request.POST.get("models_json", "[]")
    try:
        models = json.loads(raw_models)
        if not isinstance(models, list):
            raise ValueError("models must be a list")
    except (TypeError, ValueError, json.JSONDecodeError):
        return JsonResponse({"success": False, "error": "models_json must be a JSON array."}, status=400)
    try:
        memory_gb = max(0, float(request.POST.get("memory_gb", "0") or 0))
    except (TypeError, ValueError):
        memory_gb = 0
    gpu = request.POST.get("gpu", "").lower() in {"1", "true", "yes", "on"}
    return JsonResponse({
        "success": True,
        **route_model(request.POST.get("task", ""), models, memory_gb, gpu),
    })


@login_required(login_url="/login/")
@require_POST
def devcontainer_generate(request):
    project_name = request.POST.get("project_name", "Syntax Local AI")
    files = []
    raw_files = request.POST.get("files_json", "")
    if raw_files:
        try:
            parsed = json.loads(raw_files)
            if not isinstance(parsed, list) or len(parsed) > MAX_PROJECT_FILES:
                raise ValueError("invalid file list")
            files = [item for item in parsed if isinstance(item, dict)]
        except (TypeError, ValueError, json.JSONDecodeError):
            return JsonResponse({"success": False, "error": "Invalid project file list."}, status=400)
    filename = request.POST.get("filename", "")
    if filename:
        files.append({"filename": filename, "content": request.POST.get("code", "")})
    return JsonResponse({
        "success": True,
        "config": generate_devcontainer(files, project_name),
        "path": ".devcontainer/devcontainer.json",
        "notes": [
            "Review forwarded ports for your application.",
            "Keep secrets outside the container configuration.",
            "Use the same container definition in local development and CI.",
        ],
    })


@login_required(login_url="/login/")
@require_POST
def incident_analyze(request):
    return JsonResponse({
        "success": True,
        **analyze_incident(
            request.POST.get("logs", "")[:50_000],
            request.POST.get("traces", "")[:30_000],
            request.POST.get("metrics", "")[:20_000],
        ),
    })


@login_required(login_url="/login/")
@require_POST
def architecture_analyze(request):
    raw_files = request.POST.get("files_json", "")
    try:
        files = json.loads(raw_files or "[]")
        if not isinstance(files, list) or len(files) > MAX_PROJECT_FILES:
            raise ValueError("invalid architecture file list")
    except (TypeError, ValueError, json.JSONDecodeError):
        return JsonResponse({"success": False, "error": "files_json must be a JSON array."}, status=400)
    files = [
        {"filename": _safe_filename(item.get("filename", "untitled")), "content": str(item.get("content", ""))}
        for item in files if isinstance(item, dict)
    ]
    return JsonResponse({"success": True, **analyze_architecture(files)})


@login_required(login_url="/login/")
@require_POST
def cross_repository_analyze(request):
    try:
        repositories = json.loads(request.POST.get("repositories_json", "[]"))
    except json.JSONDecodeError:
        return JsonResponse({"success": False, "error": "Repositories must be valid JSON."}, status=400)
    if not isinstance(repositories, list) or not repositories:
        return JsonResponse({"success": False, "error": "Add at least one repository with files."}, status=400)
    return JsonResponse({"success": True, **analyze_cross_repository(repositories)})


@login_required(login_url="/login/")
@require_POST
def api_contract_analyze(request):
    result = analyze_api_contract(
        request.POST.get("spec", "")[:50_000],
        request.POST.get("code", "")[:50_000],
    )
    if result.get("error"):
        return JsonResponse({"success": False, "error": result["error"]}, status=400)
    return JsonResponse({"success": True, **result})


@login_required(login_url="/login/")
@require_POST
def provenance_generate(request):
    raw_files = request.POST.get("files_json", "[]")
    try:
        files = json.loads(raw_files)
        if not isinstance(files, list) or len(files) > MAX_PROJECT_FILES:
            raise ValueError("invalid provenance file list")
    except (TypeError, ValueError, json.JSONDecodeError):
        return JsonResponse({"success": False, "error": "files_json must be a JSON array."}, status=400)
    return JsonResponse({
        "success": True,
        "provenance": generate_provenance(
            [item for item in files if isinstance(item, dict)],
            request.POST.get("artifact_name", "local-artifact"),
            request.POST.get("commit", "unknown"),
        ),
    })


@login_required(login_url="/login/")
@require_POST
def provenance_verify(request):
    return JsonResponse({"success": True, **verify_provenance(request.POST.get("provenance", "")[:100_000])})


def _mcp_tools():
    return [
        {"name": "project.search", "description": "Search indexed project code.", "write": False},
        {"name": "project.read", "description": "Read one indexed project file.", "write": False},
        {"name": "project.write", "description": "Replace one indexed project file after approval.", "write": True, "approval_required": True},
        {"name": "git.status", "description": "Read local Git status.", "write": False},
        {"name": "git.diff", "description": "Read the local Git diff.", "write": False},
        {"name": "sandbox.run", "description": "Run code in the guarded sandbox.", "write": False, "approval_required": True},
        {"name": "git.stage", "description": "Stage selected local Git paths.", "write": True, "approval_required": True},
        {"name": "git.commit", "description": "Create a local Git commit.", "write": True, "approval_required": True},
    ]


def _mcp_call(user, tool_name, arguments):
    tool = next((item for item in _mcp_tools() if item["name"] == tool_name), None)
    if not tool:
        return False, {}, "Unknown MCP tool."
    if tool.get("approval_required") and not arguments.get("approved"):
        return False, {}, "This tool requires explicit approved=true."
    try:
        if tool_name == "project.search":
            return True, {"results": _search_knowledge(str(arguments.get("query", "")), owner=user)}, ""
        if tool_name == "project.read":
            filename = _safe_filename(arguments.get("filename", ""))
            document = KnowledgeDocument.objects.filter(owner=user, filename=filename, is_active=True).first()
            if not document:
                return False, {}, "Indexed file not found."
            return True, {"filename": document.filename, "content": document.original_text, "content_hash": document.content_hash}, ""
        if tool_name == "project.write":
            filename = _safe_filename(arguments.get("filename", ""))
            content = str(arguments.get("content", ""))
            if not filename or filename == "uploaded-file":
                return False, {}, "A safe project filename is required."
            if len(content.encode("utf-8")) > MAX_FILE_BYTES:
                return False, {}, "Project files are limited to 1 MB."
            existing = KnowledgeDocument.objects.filter(owner=user, filename=filename, is_active=True).first()
            expected_hash = str(arguments.get("expected_hash", "")).strip()
            if existing and expected_hash and existing.content_hash != expected_hash:
                return False, {}, "File changed since it was read; refresh before writing."
            document = _save_knowledge_document(filename, content, "manual", user)
            return True, {"filename": document.filename, "content_hash": document.content_hash, "chunks": document.chunks.count()}, ""
        if tool_name == "sandbox.run":
            result = run_sandboxed_code(
                str(arguments.get("language", "python")),
                str(arguments.get("code", "")),
            )
            return result["success"], result, result.get("stderr", "")
        if tool_name == "git.status":
            result = _run_git(["status", "--short"])
            if isinstance(result, tuple):
                return False, {}, result[1]
            return result.returncode == 0, {"raw": result.stdout}, result.stderr.strip()
        if tool_name == "git.diff":
            result = _run_git(["diff", "HEAD", "--"])
            if isinstance(result, tuple):
                return False, {}, result[1]
            return result.returncode == 0, {"diff": result.stdout[:MAX_GIT_OUTPUT_CHARS]}, result.stderr.strip()
        if tool_name == "git.stage":
            paths = arguments.get("paths") or []
            if not isinstance(paths, list) or any(not _git_path_is_safe(path) for path in paths):
                return False, {}, "Invalid Git paths."
            result = _run_git(["add", "--", *paths] if paths else ["add", "-A"])
            if isinstance(result, tuple):
                return False, {}, result[1]
            return result.returncode == 0, {"staged": paths or ["all changes"]}, result.stderr.strip()
        if tool_name == "git.commit":
            message = str(arguments.get("message", "")).strip()
            if not message or len(message) > 200:
                return False, {}, "Commit message is required and limited to 200 characters."
            result = _run_git(["commit", "-m", message])
            if isinstance(result, tuple):
                return False, {}, result[1]
            return result.returncode == 0, {"output": result.stdout}, (result.stderr or result.stdout).strip()
    except (TypeError, ValueError) as exc:
        return False, {}, str(exc)
    return False, {}, "MCP tool is not implemented."


def _mcp_log(user, tool_name, arguments, success, error):
    McpToolCall.objects.create(
        owner=user,
        tool_name=tool_name,
        arguments=arguments,
        success=success,
        error=error[:2000],
    )


def _mcp_connector_payload(connector):
    return {
        "id": connector.id,
        "name": connector.name,
        "connector_type": connector.connector_type,
        "config": connector.config,
        "enabled": connector.enabled,
        "allow_write": connector.allow_write,
    }


@login_required(login_url="/login/")
@require_http_methods(["GET", "POST"])
def mcp_connectors(request):
    if request.method == "POST":
        name = request.POST.get("name", "").strip()
        connector_type = request.POST.get("connector_type", "custom").strip()
        if not name or connector_type not in dict(McpConnector.CONNECTOR_TYPES):
            return JsonResponse({"success": False, "error": "Connector name and type are required."}, status=400)
        try:
            config = json.loads(request.POST.get("config", "{}"))
            if not isinstance(config, dict):
                raise ValueError
        except (TypeError, ValueError, json.JSONDecodeError):
            return JsonResponse({"success": False, "error": "Connector config must be a JSON object."}, status=400)
        connector = McpConnector.objects.create(
            owner=request.user,
            name=name[:120],
            connector_type=connector_type,
            config=config,
            allow_write=request.POST.get("allow_write") == "true",
        )
        return JsonResponse({"success": True, "connector": _mcp_connector_payload(connector)})
    return JsonResponse({
        "success": True,
        "connectors": [
            _mcp_connector_payload(item)
            for item in McpConnector.objects.filter(owner=request.user).order_by("name")
        ],
        "tools": _mcp_tools(),
        "recent_calls": list(
            McpToolCall.objects.filter(owner=request.user).order_by("-created_at")[:20].values(
                "tool_name", "success", "error", "created_at"
            )
        ),
    })


@login_required(login_url="/login/")
@require_http_methods(["DELETE"])
def mcp_connector_delete(request, connector_id):
    connector = get_object_or_404(McpConnector, id=connector_id, owner=request.user)
    connector.delete()
    return JsonResponse({"success": True, "deleted_id": connector_id})


@login_required(login_url="/login/")
@require_POST
def mcp_rpc(request):
    try:
        payload = json.loads(request.body or "{}")
    except (TypeError, ValueError, json.JSONDecodeError):
        return JsonResponse({"jsonrpc": "2.0", "error": {"code": -32700, "message": "Invalid JSON."}}, status=400)
    request_id = payload.get("id")
    method = payload.get("method")
    if method == "tools/list":
        result = {"tools": _mcp_tools()}
    elif method == "tools/call":
        params = payload.get("params") or {}
        tool_name = params.get("name", "")
        arguments = params.get("arguments") or {}
        success, result, error = _mcp_call(request.user, tool_name, arguments)
        _mcp_log(request.user, tool_name, mask_json(request.user, arguments), success, mask_json(request.user, error))
        if not success:
            return JsonResponse({"jsonrpc": "2.0", "id": request_id, "error": {"code": -32000, "message": mask_json(request.user, error)}}, status=400)
        result = {"content": [{"type": "text", "text": json.dumps(mask_json(request.user, result), ensure_ascii=False)}]}
    else:
        return JsonResponse({"jsonrpc": "2.0", "id": request_id, "error": {"code": -32601, "message": "Method not found."}}, status=400)
    return JsonResponse({"jsonrpc": "2.0", "id": request_id, "result": result})


@login_required(login_url="/login/")
@require_POST
def quality_analyze(request):
    code = request.POST.get("code", "")
    language = request.POST.get("language", "auto").strip().lower()
    mode = request.POST.get("mode", "all").strip().lower()
    if not code.strip():
        return JsonResponse({"success": False, "error": "Enter code to analyze."}, status=400)
    if len(code) > MAX_EXECUTION_CHARS:
        return JsonResponse({"success": False, "error": "Analysis input is too large."}, status=400)
    return JsonResponse({"success": True, **analyze_code_quality(code, language, mode)})


@login_required(login_url="/login/")
@require_POST
def security_scan(request):
    code = request.POST.get("code", "")
    filename = request.POST.get("filename", "editor-buffer")
    files = [{"filename": filename, "content": code}] if code else []
    raw_files = request.POST.get("files", "")
    if raw_files:
        try:
            submitted = json.loads(raw_files)
            if not isinstance(submitted, list) or len(submitted) > MAX_PROJECT_FILES:
                raise ValueError("invalid file list")
            files = [
                {"filename": _safe_filename(item.get("filename", "uploaded-file")), "content": str(item.get("content", ""))}
                for item in submitted
            ]
        except (TypeError, ValueError, json.JSONDecodeError):
            return JsonResponse({"success": False, "error": "Invalid security scan file list."}, status=400)
    if not files or sum(len(item["content"]) for item in files) > MAX_TEST_CHARS:
        return JsonResponse({"success": False, "error": "Provide code or a project smaller than 20,000 characters."}, status=400)
    return JsonResponse({"success": True, **scan_files(files)})


@login_required(login_url="/login/")
@require_POST
def generate_tests(request):
    code = request.POST.get("code", "")
    language = request.POST.get("language", "auto").strip().lower()
    if not code.strip():
        return JsonResponse({"success": False, "error": "Enter code before generating tests."}, status=400)
    if len(code) > MAX_TEST_CHARS:
        return JsonResponse({"success": False, "error": "Test generation input is too large."}, status=400)
    result = analyze_code_quality(code, language, "tests")
    return JsonResponse({"success": True, "test_code": result["test_template"], "summary": "Editable test scaffold generated."})


@login_required(login_url="/login/")
@require_POST
def run_tests(request):
    code = request.POST.get("code", "")
    test_code = request.POST.get("test_code", "")
    language = request.POST.get("language", "python").strip().lower()
    if not code.strip() or not test_code.strip():
        return JsonResponse({"success": False, "error": "Provide source code and test code."}, status=400)
    if len(code) + len(test_code) > MAX_TEST_CHARS:
        return JsonResponse({"success": False, "error": "Combined test input is too large."}, status=400)
    result = run_sandboxed_code(language, code.rstrip() + "\n\n" + test_code.lstrip())
    return JsonResponse({
        "success": result["success"],
        "result": result,
        "summary": "Tests passed." if result["success"] else "Tests failed.",
    })


@login_required(login_url="/login/")
@require_POST
def preview_patch(request):
    original = request.POST.get("original", "")
    updated = request.POST.get("updated", "")
    filename = request.POST.get("filename", "editor-buffer")
    diff = "".join(difflib.unified_diff(
        original.splitlines(keepends=True),
        updated.splitlines(keepends=True),
        fromfile=filename + " (current)",
        tofile=filename + " (proposed)",
        n=3,
    ))
    return JsonResponse({
        "success": True,
        "changed": original != updated,
        "diff": diff or "No changes detected.",
    })


@require_POST
def ask_code(request):
    prompt = request.POST.get("prompt", "").strip()
    language = request.POST.get("language", "auto").strip()
    code = request.POST.get("code", "").strip()
    session_id = request.POST.get("session_id", "").strip()
    edit_message_id = request.POST.get("edit_message_id", "").strip()
    requested_model = request.POST.get("model", "").strip()
    auto_route = request.POST.get("auto_route", "true").lower() in {"1", "true", "yes", "on"}
    output_format = request.POST.get("output_format", "text").strip().lower()
    output_schema_raw = request.POST.get("output_schema", "").strip()
    output_schema = None
    if output_format == "json_schema" and output_schema_raw:
        try:
            output_schema = json.loads(output_schema_raw)
            if not isinstance(output_schema, dict):
                raise ValueError
        except (TypeError, ValueError, json.JSONDecodeError):
            return JsonResponse({"success": False, "error": "Output schema must be a valid JSON object."}, status=400)
    if output_format not in {"text", "json", "json_schema"}:
        output_format = "text"
    selected_context_files = list(dict.fromkeys(
        item.strip()[:500] for item in request.POST.getlist("context_files") if item.strip()
    ))[:50]

    uploaded_files = request.FILES.getlist("files")
    image_files = request.FILES.getlist("images")
    relative_paths = request.POST.getlist("file_paths")
    single_file = request.FILES.get("file")
    if single_file and not uploaded_files:
        uploaded_files = [single_file]

    if (uploaded_files or image_files) and not request.user.is_authenticated:
        return JsonResponse({
            "success": False,
            "login_required": True,
            "error": "Create a free account or sign in to upload files and images.",
        }, status=403)

    if len(uploaded_files) > MAX_PROJECT_FILES:
        return JsonResponse({
            "success": False,
            "error": f"Please upload no more than {MAX_PROJECT_FILES} files at once.",
        }, status=400)

    uploaded_code_parts = []
    uploaded_filenames = []
    image_data = []
    image_filenames = []

    for index, uploaded_file in enumerate(uploaded_files):
        submitted_path = relative_paths[index] if index < len(relative_paths) else uploaded_file.name
        filename = _safe_filename(submitted_path)
        uploaded_filenames.append(filename)

        try:
            if uploaded_file.size > MAX_FILE_BYTES:
                uploaded_code_parts.append(
                    f"\n\n===== FILE: {filename} =====\n"
                    f"Skipped because it is larger than {MAX_FILE_BYTES // 1_000_000} MB."
                )
                continue

            raw_content = uploaded_file.read()
            file_text = _extract_document_text(filename, raw_content)

            source_type = "project" if "/" in filename else "upload"
            _save_knowledge_document(filename, file_text, source_type, request.user)
            uploaded_code_parts.append(
                f"\n\n===== FILE: {filename} =====\n{file_text}"
            )
        except Exception as ex:
            uploaded_code_parts.append(
                f"\n\n===== FILE: {filename} =====\n"
                f"Unable to read file: {str(ex)}"
            )

    for image_file in image_files:
        if image_file.size > MAX_IMAGE_BYTES:
            return JsonResponse({
                "success": False,
                "error": f"Image {image_file.name} is larger than 5 MB.",
            }, status=400)
        if not image_file.content_type.startswith("image/"):
            return JsonResponse({
                "success": False,
                "error": f"{image_file.name} is not a supported image.",
            }, status=400)
        image_filenames.append(_safe_filename(image_file.name))
        image_data.append(base64.b64encode(image_file.read()).decode("ascii"))
    uploaded_code = "\n".join(uploaded_code_parts)
    final_code = uploaded_code.strip() if uploaded_code.strip() else code

    if image_data and not prompt and not final_code:
        prompt = "Describe this image and explain any visible code or error."
    if not prompt and not final_code:
        return JsonResponse({
            "success": False,
            "error": "Please enter prompt or upload/paste code.",
        }, status=400)

    owner = request.user if request.user.is_authenticated else None
    safe_prompt = redact_sensitive_text(owner, prompt)
    safe_code = redact_sensitive_text(owner, code)
    safe_final_code = redact_sensitive_text(owner, final_code)
    user_settings = _user_ollama_settings(request.user) if owner else None
    ollama_base_url = user_settings.server_url if user_settings else OLLAMA_BASE_URL
    configured_default_model = user_settings.default_model if user_settings else DEFAULT_MODEL
    routing_reason = "manual model selection"
    if auto_route and not requested_model:
        task_for_routing = " ".join([prompt, code, "image" if image_data else ""]).strip()
        try:
            hardware_memory = max(0, float(os.environ.get("LOCAL_AI_MEMORY_GB", "0") or 0))
        except (TypeError, ValueError):
            hardware_memory = 0
        hardware_gpu = os.environ.get("LOCAL_AI_GPU", "").lower() in {"1", "true", "yes", "on"}
        routed = route_model(task_for_routing, _available_models(ollama_base_url), hardware_memory, hardware_gpu)
        configured_default_model = routed["selected_model"]
        routing_reason = routed["reason"]
    revision_branch_id = None
    if session_id:
        session = get_object_or_404(ChatSession, id=session_id, owner=owner)
        model_name = requested_model or session.model_name or configured_default_model
        if image_data and model_name == DEFAULT_MODEL:
            model_name = VISION_MODEL
        if session.model_name != model_name:
            session.model_name = model_name
            session.save(update_fields=["model_name", "updated_at"])
    else:
        title = prompt[:60] if prompt else (
            uploaded_filenames[0][:60] if uploaded_filenames else "New Chat"
        )
        model_name = requested_model or configured_default_model
        if image_data and model_name == DEFAULT_MODEL:
            model_name = VISION_MODEL
        session = ChatSession.objects.create(
            owner=owner,
            title=title,
            model_name=model_name,
        )

    if edit_message_id:
        if not session_id:
            return JsonResponse({"success": False, "error": "An existing chat is required to edit a message."}, status=400)
        target = get_object_or_404(
            ChatMessage,
            id=edit_message_id,
            session=session,
            role="user",
        )
        revision_branch_id = uuid.uuid4()
        previous_messages = ChatMessage.objects.filter(
            session=session,
            id__gte=target.id,
        ).order_by("id")
        ConversationRevision.objects.bulk_create([
            ConversationRevision(
                session=session,
                source_message_id=message.id,
                branch_id=revision_branch_id,
                role=message.role,
                content=message.content,
                model_name=message.model_name,
            )
            for message in previous_messages
        ])
        previous_messages.delete()

    user_message_parts = []
    if prompt:
        user_message_parts.append(f"Prompt:\n{safe_prompt}")
    if uploaded_filenames:
        user_message_parts.append("Uploaded Files:\n" + "\n".join(uploaded_filenames))
    if selected_context_files:
        user_message_parts.append("Selected Project Files:" + chr(10) + chr(10).join(selected_context_files))
    if image_filenames:
        user_message_parts.append("Uploaded Images:\n" + "\n".join(image_filenames))
    if final_code:
        user_message_parts.append(f"Code:\n{safe_final_code[:6000]}")

    user_message = ChatMessage.objects.create(
        session=session,
        role="user",
        content="\n\n".join(user_message_parts),
        filename=", ".join(uploaded_filenames + image_filenames),
        model_name=model_name,
    )

    max_code_chars = 24_000
    truncated_note = ""
    if len(final_code) > max_code_chars:
        final_code = final_code[:max_code_chars]
        truncated_note = "\n\nNote: the uploaded code was truncated for the local model."
    safe_final_code = redact_sensitive_text(owner, final_code)

    relevant_chunks = _search_knowledge(
        safe_prompt or safe_code,
        owner=owner,
        filenames=selected_context_files or None,
    )
    knowledge_context = "\n\n".join(
        f"===== PROJECT CONTEXT: {item['filename']}:{item['line_start']}-{item['line_end']} =====\n"
        f"{item['content']}"
        for item in relevant_chunks
    )

    full_prompt = f"""
You are a fully offline coding assistant running locally.

Language:
{language}

User request:
 {safe_prompt or "(No separate prompt was provided. Interpret the Code section as the user's request if it contains natural language.)"}

Uploaded files (optional):
{", ".join(uploaded_filenames) or "None"}

Relevant project context (optional):
{knowledge_context or "None"}

Code or additional user input (optional):
 {safe_final_code or "None"}

{truncated_note}

Output format:
{("Return valid JSON only." if output_format == "json" else "Return JSON matching this schema exactly: " + json.dumps(output_schema) if output_format == "json_schema" else "Use normal Markdown/text.")}

Instructions:
1. Answer the user's request directly, even when no files, project context, or source code are provided.
2. Treat uploaded files and project context as optional supporting information, not a prerequisite for an answer.
3. Never ask the user to upload or provide files merely because none were supplied.
4. If the user asks for code, a question, an example, or an explanation, provide it in this response.
5. If there is an error, show the exact problem and give corrected code where appropriate.
6. Mention a file name only when a file was actually provided or a file is needed for a concrete change.
7. Keep the answer practical and do not say you need internet.
"""
    ai_event = None
    event_started = time.perf_counter()
    try:
        ai_event = AiEvent.objects.create(
            owner=owner,
            event_type="chat",
            model_name=model_name,
            input_chars=len(full_prompt),
        )
        fallback_model = user_settings.fallback_model if user_settings else ""
        candidate_models = list(dict.fromkeys(
            candidate for candidate in (model_name, fallback_model) if candidate
        ))
        response = None
        last_error = None
        for candidate_model in candidate_models:
            try:
                candidate_response = requests.post(
                    f"{ollama_base_url}/api/generate",
                    json={
                        "model": candidate_model,
                        "prompt": full_prompt,
                        "stream": True,
                        "options": {
                            "temperature": user_settings.temperature if user_settings else 0.2,
                            "top_p": user_settings.top_p if user_settings else 0.9,
                            "num_ctx": max(512, (user_settings.max_context_chars if user_settings else 24000) // 4),
                        },
                        **({"format": output_schema or "json"} if output_format in {"json", "json_schema"} else {}),
                        **({"images": image_data} if image_data else {}),
                    },
                    timeout=(10, 300),
                    stream=True,
                )
            except requests.RequestException as exc:
                last_error = exc
                continue
            response = candidate_response
            if response.status_code == 200:
                model_name = candidate_model
                break

        if response is None:
            raise last_error or requests.exceptions.ConnectionError("No Ollama model responded.")
        if ai_event.model_name != model_name:
            ai_event.model_name = model_name
            ai_event.save(update_fields=["model_name"])
        if response.status_code != 200:
            ai_event.success = False
            ai_event.duration_ms = round((time.perf_counter() - event_started) * 1000)
            ai_event.metadata = {"status_code": response.status_code, "error": response.text[:500]}
            ai_event.save(update_fields=["success", "duration_ms", "metadata"])
            return JsonResponse({"success": False, "error": response.text}, status=500)

        def stream_answer():
            answer_parts = []
            first_token_ms = None
            try:
                for raw_line in response.iter_lines(decode_unicode=True):
                    if not raw_line:
                        continue
                    data = json.loads(raw_line)
                    token = data.get("response", "")
                    if token:
                        if first_token_ms is None:
                            first_token_ms = round((time.perf_counter() - event_started) * 1000)
                        answer_parts.append(token)
                        yield _event({"type": "token", "token": token})
                    if data.get("done"):
                        break

                answer = "".join(answer_parts).strip()
                ChatMessage.objects.create(
                    session=session,
                    role="assistant",
                    content=answer,
                    model_name=model_name,
                )
                ai_event.duration_ms = round((time.perf_counter() - event_started) * 1000)
                ai_event.output_chars = len(answer)
                ai_event.metadata = {
                    "sources": len(relevant_chunks),
                    "images": len(image_data),
                    "session_id": session.id,
                    "context_chars": len(full_prompt),
                    "output_chars_per_second": round(len(answer) / max(ai_event.duration_ms / 1000, 0.001), 1),
                    "first_token_ms": first_token_ms or 0,
                    "output_format": output_format,
                    "routing_reason": routing_reason,
                }
                ai_event.save(update_fields=["duration_ms", "output_chars", "metadata"])
                session.save(update_fields=["updated_at"])
                yield _event({
                    "type": "complete",
                    "session_id": session.id,
                    "title": session.title,
                    "model": model_name,
                    "user_message_id": user_message.id,
                    "revision_branch_id": str(revision_branch_id) if revision_branch_id else "",
                    "sources": [
                        {
                            "filename": item["filename"],
                            "language": item["language"],
                            "line_start": item["line_start"],
                            "line_end": item["line_end"],
                            "page_start": item.get("page_start"),
                            "page_end": item.get("page_end"),
                            "content": item["content"],
                        }
                        for item in relevant_chunks
                    ],
                    "answer": answer,
                })
            except Exception as ex:
                if ai_event:
                    ai_event.success = False
                    ai_event.duration_ms = round((time.perf_counter() - event_started) * 1000)
                    ai_event.metadata = {"error": str(ex)[:500]}
                    ai_event.save(update_fields=["success", "duration_ms", "metadata"])
                yield _event({"type": "error", "error": str(ex)})

        streaming_response = StreamingHttpResponse(
            stream_answer(),
            content_type="application/x-ndjson",
        )
        streaming_response["Cache-Control"] = "no-cache"
        return streaming_response
    except requests.exceptions.ConnectionError:
        if ai_event:
            ai_event.success = False
            ai_event.duration_ms = round((time.perf_counter() - event_started) * 1000)
            ai_event.metadata = {"error": "Ollama is not running"}
            ai_event.save(update_fields=["success", "duration_ms", "metadata"])
        return JsonResponse({
            "success": False,
            "error": "Ollama is not running. Start Ollama first.",
        }, status=500)
    except requests.exceptions.Timeout:
        if ai_event:
            ai_event.success = False
            ai_event.duration_ms = round((time.perf_counter() - event_started) * 1000)
            ai_event.metadata = {"error": "Model response timeout"}
            ai_event.save(update_fields=["success", "duration_ms", "metadata"])
        return JsonResponse({
            "success": False,
            "error": "Model response timeout. Try smaller files or a smaller prompt.",
        }, status=500)
    except Exception as ex:
        if ai_event:
            ai_event.success = False
            ai_event.duration_ms = round((time.perf_counter() - event_started) * 1000)
            ai_event.metadata = {"error": str(ex)[:500]}
            ai_event.save(update_fields=["success", "duration_ms", "metadata"])
        return JsonResponse({"success": False, "error": str(ex)}, status=500)
